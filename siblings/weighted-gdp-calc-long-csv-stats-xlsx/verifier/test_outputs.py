import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook

W = Path("/root/work")
OUT = W / "nx_stats.xlsx"
LABELS = ["Min", "Max", "Median", "Mean", "P25", "P75", "Weighted Mean"]
TOL = 0.002


def pct(vals, q):
    s = sorted(vals)
    h = (len(s) - 1) * q
    lo = int(h)
    hi = min(lo + 1, len(s) - 1)
    return s[lo] + (h - lo) * (s[hi] - s[lo])


def expected():
    d = {}
    with open(W / "trade_long.csv") as f:
        for r in csv.DictReader(f):
            v = float(r["value"])
            if r["unit"].strip() == "USD mn":
                v /= 1000.0
            d[(r["country"], int(r["year"]), r["indicator"])] = v
    years = sorted({k[1] for k in d})
    countries = sorted({k[0] for k in d})
    res = {}
    for y in years:
        nx, g = [], []
        for c in countries:
            gd = d[(c, y, "gdp_current")]
            nx.append((d[(c, y, "exports_gs")] - d[(c, y, "imports_gs")]) / gd * 100)
            g.append(gd)
        res[y] = [min(nx), max(nx), pct(nx, 0.5), sum(nx) / len(nx), pct(nx, 0.25),
                  pct(nx, 0.75), sum(a * b for a, b in zip(nx, g)) / sum(g)]
    return years, res


@pytest.fixture(scope="module")
def ws():
    assert OUT.exists(), "output missing"
    wb = load_workbook(OUT, data_only=True)
    assert "Stats" in wb.sheetnames
    return wb["Stats"]


def test_file_exists():
    assert OUT.exists()


def test_header(ws):
    years, _ = expected()
    assert str(ws["A1"].value).strip() == "Statistic"
    got = [ws.cell(1, 2 + i).value for i in range(len(years))]
    assert [int(float(v)) for v in got] == years


def test_labels(ws):
    assert [str(ws.cell(2 + i, 1).value).strip() for i in range(len(LABELS))] == LABELS


def test_numeric(ws):
    years, _ = expected()
    for i in range(len(LABELS)):
        for j in range(len(years)):
            v = ws.cell(2 + i, 2 + j).value
            assert isinstance(v, (int, float)), (i, j, v)


def test_values(ws):
    years, res = expected()
    for j, y in enumerate(years):
        for i, lab in enumerate(LABELS):
            v = ws.cell(2 + i, 2 + j).value
            assert abs(float(v) - res[y][i]) <= TOL, (lab, y, v, res[y][i])
