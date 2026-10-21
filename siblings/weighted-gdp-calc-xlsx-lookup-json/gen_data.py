import random
from openpyxl import Workbook

random.seed(4242)
C = [("Bahrain", "419", 38000), ("Kuwait", "443", 140000), ("Oman", "449", 80000),
     ("Qatar", "453", 180000), ("Saudi Arabia", "456", 800000), ("United Arab Emirates", "466", 420000)]
years = list(range(2018, 2024))
rows = []
for name, code, base in C:
    g = [round(base * random.uniform(0.85, 1.25), 1) for _ in years]
    x = [round(v * random.uniform(0.35, 0.95), 1) for v in g]
    m = [round(v * random.uniform(0.3, 0.8), 1) for v in g]
    gx = [round(v * random.uniform(0.2, 0.4), 1) for v in g]
    rows.append([name, code + "NGDPD.A", "Gross Domestic Product, Current Prices"] + g)
    rows.append([name, code + "BXGS_BP6.A", "Exports, Goods and Services"] + x)
    rows.append([name, code + "BMGS_BP6.A", "Imports, Goods and Services"] + m)
    rows.append([name, code + "GGX.A", "General Government Expenditure"] + gx)
random.shuffle(rows)

wb = Workbook()
notes = wb.active
notes.title = "Notes"
notes["A1"] = "Synthetic GCC dataset. See sheet Data."
ws = wb.create_sheet("Data")
ws.cell(1, 1, "GCC external sector (USD millions)")
ws.cell(2, 1, "Source: synthetic IMF-style extract")
hdr = ["Country", "Series Code", "Series Name"] + years
for j, v in enumerate(hdr, 1):
    ws.cell(3, j, v)
for i, r in enumerate(rows, 4):
    for j, v in enumerate(r, 1):
        ws.cell(i, j, v)
wb.save("gdp.xlsx")
