import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

W = Path("/root/work")
OUT = W / "nx_gdp.json"
TOL = 0.011


def load_data():
    wb = load_workbook(W / "gdp.xlsx", data_only=True)
    ws = wb["Data"]
    hdr = None
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, 2).value == "Series Code":
            hdr = r
            break
    ycols = {}
    for c in range(4, ws.max_column + 1):
        v = ws.cell(hdr, c).value
        if v is not None:
            ycols[int(v)] = c
    data, names = {}, {}
    for r in range(hdr + 1, ws.max_row + 1):
        code = ws.cell(r, 2).value
        if not code:
            continue
        data[code] = {y: float(ws.cell(r, c).value) for y, c in ycols.items()}
        names[code[:3]] = ws.cell(r, 1).value
    return data, names, sorted(ycols)


def expected():
    data, names, years = load_data()
    ctry, wt = {}, {}
    for y in years:
        num = den = 0.0
        for p, n in names.items():
            g = data[p + "NGDPD.A"][y]
            nx = (data[p + "BXGS_BP6.A"][y] - data[p + "BMGS_BP6.A"][y]) / g * 100
            ctry.setdefault(n, {})[str(y)] = nx
            num += nx * g
            den += g
        wt[str(y)] = num / den
    return ctry, wt


@pytest.fixture(scope="module")
def out():
    assert OUT.exists(), "output file missing"
    return json.loads(OUT.read_text())


def test_file_exists():
    assert OUT.exists()


def test_schema(out):
    ctry, wt = expected()
    assert set(out.keys()) >= {"countries", "gcc_weighted"}
    assert set(out["countries"].keys()) == set(ctry.keys())
    for c in ctry:
        assert set(out["countries"][c].keys()) == set(ctry[c].keys())
    assert set(out["gcc_weighted"].keys()) == set(wt.keys())


def test_country_values(out):
    ctry, _ = expected()
    for c, ys in ctry.items():
        for y, v in ys.items():
            assert abs(float(out["countries"][c][y]) - v) <= TOL, (c, y, out["countries"][c][y], v)


def test_weighted_values(out):
    _, wt = expected()
    for y, v in wt.items():
        assert abs(float(out["gcc_weighted"][y]) - v) <= TOL, (y, out["gcc_weighted"][y], v)
