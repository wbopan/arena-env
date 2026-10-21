import os
import random
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from docx import Document
from pptx import Presentation
from pptx.util import Inches

random.seed(20240611)
KW = {
    'LLM': ['transformer', 'token', 'prompt', 'finetuning', 'attention'],
    'trapped_ion_and_qc': ['ion', 'qubit', 'laser', 'entanglement', 'trap'],
    'black_hole': ['horizon', 'accretion', 'schwarzschild', 'hawking', 'singularity'],
    'DNA': ['nucleotide', 'helix', 'genome', 'polymerase', 'sequencing'],
    'music_history': ['baroque', 'symphony', 'composer', 'opera', 'renaissance'],
}
FILLER = ('the study results method data analysis we report section figure table model '
          'sample review notes summary approach measured observed context overview').split()
SUBJ = list(KW)
os.makedirs('inbox', exist_ok=True)


def make_words(primary):
    ws = [random.choice(FILLER) for _ in range(random.randint(35, 60))]
    for _ in range(random.randint(8, 12)):
        ws.append(random.choice(KW[primary]))
    others = [s for s in SUBJ if s != primary]
    for _ in range(random.randint(0, 3)):
        ws.append(random.choice(KW[random.choice(others)]))
    random.shuffle(ws)
    return ws


def split(ws, k):
    return [ws[i::k] for i in range(k)]


subjects = SUBJ * 3
random.shuffle(subjects)
types = ['pdf'] * 6 + ['docx'] * 5 + ['pptx'] * 4
random.shuffle(types)
prefixes = ['scan', 'doc', 'file', 'misc', 'notes', 'draft']
used = set()
for subj, ext in zip(subjects, types):
    while True:
        name = f"{random.choice(prefixes)}_{random.randint(100, 999)}.{ext}"
        if name not in used:
            used.add(name)
            break
    path = os.path.join('inbox', name)
    ws = make_words(subj)
    if ext == 'pdf':
        chunks = split(ws, random.randint(1, 3))
        c = canvas.Canvas(path, pagesize=letter)
        for ch in chunks:
            y = 720
            for i in range(0, len(ch), 10):
                c.drawString(72, y, ' '.join(ch[i:i + 10]))
                y -= 18
            c.showPage()
        c.save()
    elif ext == 'docx':
        chunks = split(ws, random.randint(3, 6))
        d = Document()
        for j, ch in enumerate(chunks):
            d.add_paragraph(' '.join(ch))
            if j == 0:
                d.add_paragraph('')
                d.add_paragraph('   ')
        d.save(path)
    else:
        chunks = split(ws, random.randint(2, 4))
        p = Presentation()
        for ch in chunks:
            s = p.slides.add_slide(p.slide_layouts[6])
            tb = s.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(5))
            tb.text_frame.text = ' '.join(ch)
        p.save(path)
