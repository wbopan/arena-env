import csv
import os
import re
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from pypdf import PdfReader

WORK = Path('/root/work')
INBOX = WORK / 'inbox'
OUT = WORK / 'manifest.csv'
KW = {
    'LLM': ['transformer', 'token', 'prompt', 'finetuning', 'attention'],
    'trapped_ion_and_qc': ['ion', 'qubit', 'laser', 'entanglement', 'trap'],
    'black_hole': ['horizon', 'accretion', 'schwarzschild', 'hawking', 'singularity'],
    'DNA': ['nucleotide', 'helix', 'genome', 'polymerase', 'sequencing'],
    'music_history': ['baroque', 'symphony', 'composer', 'opera', 'renaissance'],
}


def analyse(p):
    ext = p.suffix.lower()
    if ext == '.pdf':
        r = PdfReader(str(p))
        text = '\n'.join((pg.extract_text() or '') for pg in r.pages)
        units = len(r.pages)
    elif ext == '.docx':
        paras = [x.text for x in Document(str(p)).paragraphs]
        text = '\n'.join(paras)
        units = sum(1 for t in paras if t.strip())
    else:
        prs = Presentation(str(p))
        parts = []
        for s in prs.slides:
            for sh in s.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
        text = '\n'.join(parts)
        units = len(prs.slides)
    toks = re.findall(r'[a-z]+', text.lower())
    scores = {s: sum(toks.count(k) for k in kws) for s, kws in KW.items()}
    best = sorted(scores, key=lambda s: (-scores[s], s))[0]
    return best, units


def expected():
    files = sorted(f for f in os.listdir(INBOX) if f.lower().endswith(('.pdf', '.docx', '.pptx')))
    return {f: analyse(INBOX / f) for f in files}


def rows():
    with open(OUT, newline='') as fh:
        return list(csv.DictReader(fh))


def test_exists():
    assert OUT.is_file()


def test_header():
    with open(OUT, newline='') as fh:
        header = next(csv.reader(fh))
    assert [h.strip() for h in header] == ['file', 'subject', 'units']


def test_rows_cover_inbox_sorted():
    exp = expected()
    got = [r['file'].strip() for r in rows()]
    assert got == sorted(exp.keys())


def test_subjects():
    exp = expected()
    got = {r['file'].strip(): r['subject'].strip() for r in rows()}
    for f, (s, _) in exp.items():
        assert got.get(f) == s, f


def test_units():
    exp = expected()
    got = {r['file'].strip(): r['units'].strip() for r in rows()}
    for f, (_, u) in exp.items():
        assert int(float(got[f])) == u, f
