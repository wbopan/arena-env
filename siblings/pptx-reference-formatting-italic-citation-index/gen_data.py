import random
import re
import zipfile
from pptx import Presentation
from pptx.util import Inches, Pt

random.seed(1729)
PAPERS = [
    'ReAct: Synergizing Reasoning and Acting in Language Models',
    'Toolformer: Language Models Can Teach Themselves to Use Tools',
    'Reflexion: Language Agents with Verbal Reinforcement Learning',
    'Voyager: An Open-Ended Embodied Agent with Large Language Models',
    'Generative Agents: Interactive Simulacra of Human Behavior',
    'AutoGen: Enabling Next-Gen LLM Applications via Multi-Agent Conversation',
    'Why Do Multi-Agent LLM Systems Fail?',
    'SWE-agent: Agent-Computer Interfaces Enable Automated Software Engineering',
    'Tree of Thoughts: Deliberate Problem Solving with Large Language Models',
    'WebArena: A Realistic Web Environment for Building Autonomous Agents',
]
TOPICS = ['Planning', 'Tool Use', 'Memory', 'Multi-Agent Coordination', 'Evaluation', 'Self-Reflection', 'Web Agents', 'Coding Agents']
picked = random.sample(PAPERS, 7)


def variant(t):
    k = random.randint(0, 2)
    if k == 0:
        return t.upper()
    if k == 1:
        w = t.split(' ')
        i = random.randint(1, len(w) - 1)
        return ' '.join(w[:i]) + '   ' + ' '.join(w[i:])
    return '  ' + t.lower() + ' '


def box(s, runs):
    b = s.shapes.add_textbox(Inches(random.uniform(0.3, 3.0)), Inches(random.uniform(4.6, 6.6)), Inches(6), Inches(0.5))
    p = b.text_frame.paragraphs[0]
    for text, it in runs:
        r = p.add_run()
        r.text = text
        r.font.italic = it
        r.font.size = Pt(random.choice([11, 12, 14]))


def shrink(path):
    with zipfile.ZipFile(path) as z:
        items = [(n, z.read(n)) for n in z.namelist()]
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for n, d in items:
            if 'thumbnail' in n:
                continue
            if n == '_rels/.rels':
                d = re.sub(rb'<Relationship[^>]*thumbnail[^>]*/>', b'', d)
            z.writestr(n, d)


prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = 'Agent Papers Reading Group'
s.placeholders[1].text = 'Quarterly digest'
seen = []
for n in range(2, 15):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = f'{random.choice(TOPICS)} - part {n}'
    tf = s.placeholders[1].text_frame
    tf.text = f'Key finding {random.randint(10, 99)}'
    tf.add_paragraph().text = f'Benchmark gain of {random.randint(2, 40)} points'
    if random.random() < 0.3:
        r = tf.add_paragraph().add_run()
        r.text = random.choice(PAPERS)
        r.font.italic = True
    for _ in range(random.choice([0, 1, 1, 2])):
        base = random.choice(picked)
        t = variant(base) if (base in seen and random.random() < 0.6) else base
        seen.append(base)
        if random.random() < 0.35:
            w = t.split(' ')
            i = len(w) // 2
            box(s, [(' '.join(w[:i]) + ' ', True), (' '.join(w[i:]), True)])
        else:
            box(s, [(t, True)])
    d = random.random()
    if d < 0.3:
        box(s, [('Source: internal benchmark logs', False)])
    elif d < 0.55:
        box(s, [('See also ', False), (random.choice(PAPERS), True)])
prs.save('slides.pptx')
shrink('slides.pptx')
