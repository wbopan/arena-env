import json
from pathlib import Path

import pytest
from pptx import Presentation

IN = Path('/root/work/slides.pptx')
OUT = Path('/root/work/references.json')


def norm(s):
    return ' '.join(s.split())


def expected():
    prs = Presentation(str(IN))
    order, slides, first = [], {}, {}
    for sn, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.is_placeholder or not shape.has_text_frame:
                continue
            paras = shape.text_frame.paragraphs
            textual = [r for p in paras for r in p.runs if r.text.strip()]
            if not textual:
                continue
            if not all(r.font.italic is True for r in textual):
                continue
            text = norm(' '.join(''.join(r.text for r in p.runs) for p in paras))
            key = text.lower()
            if key not in first:
                first[key] = text
                order.append(key)
                slides[key] = set()
            slides[key].add(sn)
    return [{'number': i + 1, 'title': first[k], 'slides': sorted(slides[k])} for i, k in enumerate(order)]


@pytest.fixture(scope='module')
def out():
    assert OUT.exists(), 'references.json missing'
    return json.loads(OUT.read_text())


@pytest.fixture(scope='module')
def exp():
    return expected()


def test_file_exists():
    assert OUT.exists()
    json.loads(OUT.read_text())


def test_schema(out):
    assert isinstance(out, dict)
    refs = out.get('references')
    assert isinstance(refs, list)
    for i, r in enumerate(refs, 1):
        assert isinstance(r, dict)
        assert {'number', 'title', 'slides'} <= set(r)
        assert r['number'] == i
        assert isinstance(r['title'], str)
        assert isinstance(r['slides'], list)
        assert all(isinstance(x, int) for x in r['slides'])


def test_count(out, exp):
    assert len(out['references']) == len(exp)


def test_titles_and_order(out, exp):
    assert [norm(r['title']) for r in out['references']] == [e['title'] for e in exp]


def test_slide_lists(out, exp):
    got = {norm(r['title']).lower(): list(r['slides']) for r in out['references']}
    for e in exp:
        assert got.get(e['title'].lower()) == e['slides'], e['title']
