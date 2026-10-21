import json
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

W = Path("/root/work")
OUT = W / "dept_summary.json"
KEYS = ["department", "headcount_old", "headcount_new", "payroll_old", "payroll_new", "payroll_change"]


def load_old():
    d = {}
    for page in PdfReader(str(W / "roster_backup.pdf")).pages:
        for line in (page.extract_text() or "").splitlines():
            line = line.strip()
            if re.match(r"^EMP\d{5}\s*\|", line):
                p = [x.strip() for x in line.split("|")]
                d[p[0]] = (p[2], int(re.sub(r"[^0-9]", "", p[3])))
    return d


def load_new():
    ws = load_workbook(W / "roster_current.xlsx", data_only=True)["Employees"]
    rows = list(ws.iter_rows(values_only=True))
    h = list(rows[0])
    i, j, k = h.index("Employee ID"), h.index("Department"), h.index("Annual Salary")
    return {str(r[i]).strip(): (str(r[j]).strip(), int(r[k])) for r in rows[1:] if r[i] is not None}


def expected():
    o, n = load_old(), load_new()
    depts = sorted({v[0] for v in o.values()} | {v[0] for v in n.values()})
    rows = []
    for d in depts:
        ho = sum(1 for v in o.values() if v[0] == d)
        hn = sum(1 for v in n.values() if v[0] == d)
        po = sum(v[1] for v in o.values() if v[0] == d)
        pn = sum(v[1] for v in n.values() if v[0] == d)
        rows.append({"department": d, "headcount_old": ho, "headcount_new": hn,
                     "payroll_old": po, "payroll_new": pn, "payroll_change": pn - po})
    transfers = sum(1 for k in o if k in n and o[k][0] != n[k][0])
    return rows, transfers, len(n) - len(o)


@pytest.fixture(scope="module")
def out():
    assert OUT.exists(), f"{OUT} not found"
    with open(OUT) as f:
        return json.load(f)


def test_output_exists_and_json():
    assert OUT.exists(), f"{OUT} not found"
    with open(OUT) as f:
        json.load(f)


def test_schema(out):
    assert isinstance(out.get("departments"), list)
    assert "transfers" in out and "net_headcount_change" in out
    for d in out["departments"]:
        for k in KEYS:
            assert k in d, f"missing key {k} in {d}"
        for k in KEYS[1:]:
            assert isinstance(d[k], (int, float)) and not isinstance(d[k], bool), f"{k} not numeric in {d}"


def test_department_names_and_order(out):
    exp, _, _ = expected()
    got = [d["department"] for d in out["departments"]]
    assert got == [e["department"] for e in exp], f"got {got}"


def test_department_values(out):
    exp, _, _ = expected()
    got = {d["department"]: d for d in out["departments"]}
    for e in exp:
        g = got.get(e["department"])
        assert g is not None, f"{e['department']} missing"
        for k in KEYS[1:]:
            assert g[k] == e[k], f"{e['department']} {k}: got {g[k]}, expected {e[k]}"


def test_transfers(out):
    _, t, _ = expected()
    assert out["transfers"] == t, f"transfers got {out['transfers']}, expected {t}"


def test_net_headcount_change(out):
    _, _, n = expected()
    assert out["net_headcount_change"] == n, f"got {out['net_headcount_change']}, expected {n}"
