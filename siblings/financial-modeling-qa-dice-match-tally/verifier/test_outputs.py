import json
import os

from openpyxl import load_workbook

DATA = "/root/work/data.xlsx"
OUT = "/root/work/answer.json"
KEYS = ["player1_wins", "player2_wins", "ties", "difference"]


def game_score(dice):
    a, b, c = sorted(dice)
    if a == c:
        return 300 + a
    if b == a + 1 and c == b + 1:
        return 200 + c
    if a == b:
        return 100 + 10 * a + c
    if b == c:
        return 100 + 10 * b + a
    return a + b + c


def expected():
    ws = load_workbook(DATA, data_only=True)["Games"]
    scores = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue
        scores[int(row[0])] = game_score([int(row[1]), int(row[2]), int(row[3])])
    p1 = p2 = t = 0
    for g in range(1, max(scores) + 1, 2):
        s1, s2 = scores[g], scores[g + 1]
        if s1 > s2:
            p1 += 1
        elif s2 > s1:
            p2 += 1
        else:
            t += 1
    return {"player1_wins": p1, "player2_wins": p2, "ties": t, "difference": p1 - p2}


def load():
    with open(OUT) as f:
        return json.load(f)


def test_output_exists():
    assert os.path.exists(OUT), "answer.json missing"


def test_schema():
    d = load()
    assert isinstance(d, dict)
    for k in KEYS:
        assert k in d, f"missing key {k}"
        assert isinstance(d[k], int) and not isinstance(d[k], bool), f"{k} must be int"


def test_counts_consistent():
    d = load()
    assert d["player1_wins"] + d["player2_wins"] + d["ties"] == 100
    assert d["difference"] == d["player1_wins"] - d["player2_wins"]


def test_values_correct():
    d = load()
    e = expected()
    for k in KEYS:
        assert d[k] == e[k], f"{k}: expected {e[k]}, got {d[k]}"
