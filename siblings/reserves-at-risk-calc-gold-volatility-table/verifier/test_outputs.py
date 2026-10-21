import math
from pathlib import Path

import numpy as np
import pytest
from openpyxl import load_workbook

W = Path("/root/work")
OUT = W / "gold_vol.xlsx"
TOL = 1e-5
HEADER = ["Period", "Price", "LogReturn", "Vol3M", "Vol12M", "Vol3M_Ann"]


def inputs():
    ws = load_workbook(W / "gold_prices.xlsx", data_only=True)["Prices"]
    per, pr = [], []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[0] is None:
            continue
        per.append(str(row[0]).strip())
        pr.append(float(row[1]))
    return per, np.array(pr)


def expected():
    per, pr = inputs()
    r = np.concatenate([[np.nan], 100 * np.diff(np.log(pr))])
    n = len(pr)
    v3 = np.full(n, np.nan)
    v12 = np.full(n, np.nan)
    for i in range(n):
        if i >= 3:
            v3[i] = np.std(r[i - 2:i + 1], ddof=1)
        if i >= 12:
            v12[i] = np.std(r[i - 11:i + 1], ddof=1)
    return per, pr, r, v3, v12, v3 * math.sqrt(12)


def out_rows():
    ws = load_workbook(OUT, data_only=True)["Volatility"]
    rows = [list(r) for r in ws.iter_rows(min_row=1, values_only=True)]
    hdr = [str(c).strip() if c is not None else None for c in rows[0][:6]]
    data = [r for r in rows[1:] if r and r[0] is not None]
    return hdr, data


def check_col(idx, exp):
    _, data = out_rows()
    for k, e in enumerate(exp):
        got = data[k][idx] if len(data[k]) > idx else None
        if np.isnan(e):
            assert got is None or got == "", f"row {k+2} col {idx}: expected blank, got {got}"
        else:
            assert isinstance(got, (int, float)), f"row {k+2} col {idx}: not numeric: {got!r}"
            assert abs(got - e) <= TOL, f"row {k+2} col {idx}: {got} vs {e}"


def test_file_exists():
    assert OUT.exists()
    wb = load_workbook(OUT)
    assert "Volatility" in wb.sheetnames


def test_header_and_periods():
    hdr, data = out_rows()
    assert hdr == HEADER
    per, pr, *_ = expected()
    assert len(data) == len(per)
    assert [str(d[0]).strip() for d in data] == per
    for d, p in zip(data, pr):
        assert abs(float(d[1]) - p) < 1e-6


def test_log_returns():
    _, _, r, _, _, _ = expected()
    check_col(2, r)


def test_vol3m_and_annualized():
    _, _, _, v3, _, ann = expected()
    check_col(3, v3)
    check_col(5, ann)


def test_vol12m():
    _, _, _, _, v12, _ = expected()
    check_col(4, v12)
