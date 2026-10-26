from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

ROOT = Path('/root/work')
SOURCE = ROOT / 'expression_panel.xlsx'
OUTPUT = ROOT / 'sample_qc.xlsx'

def books():
    return load_workbook(SOURCE), load_workbook(OUTPUT)

def test_file_exists():
    assert OUTPUT.is_file()

def test_sheets_and_preserved_inputs():
    original, result = books()
    assert result.sheetnames == ['Data', 'Metadata', 'QC']
    for name in ('Data', 'Metadata'):
        a, b = original[name], result[name]
        assert a.max_row == b.max_row and a.max_column == b.max_column
        assert list(a.values) == list(b.values)

def test_qc_schema_and_order():
    original, result = books()
    qc = result['QC']
    metadata = list(original['Metadata'].values)[1:]
    assert [qc.cell(1, col).value for col in range(1, 6)] == ['Sample_ID', 'Group', 'Detected_Proteins', 'Mean_Log2', 'SD_Log2']
    assert qc.max_row == len(metadata) + 1 and qc.max_column == 5
    assert [(qc.cell(r, 1).value, qc.cell(r, 2).value) for r in range(2, qc.max_row + 1)] == metadata

def test_formulas_reference_correct_full_columns():
    original, result = books()
    data = original['Data']
    positions = {data.cell(1, c).value: c for c in range(2, data.max_column + 1)}
    qc = result['QC']
    for r in range(2, qc.max_row + 1):
        column = get_column_letter(positions[qc.cell(r, 1).value])
        reference = f'Data!{column}2:{column}{data.max_row}'
        assert [qc.cell(r, c).value for c in (3, 4, 5)] == [f'=COUNT({reference})', f'=AVERAGE({reference})', f'=STDEV({reference})']

def test_display_format():
    _, result = books()
    qc = result['QC']
    assert all(qc.cell(r, c).number_format == '0.0000' for r in range(2, qc.max_row + 1) for c in (4, 5))
