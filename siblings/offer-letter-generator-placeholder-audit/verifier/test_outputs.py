import json
import re
from pathlib import Path

import pytest
from docx import Document

WORK = Path("/root/work")
TEMPLATE = WORK / "offer_letter_template.docx"
DATA = WORK / "employee_data.json"
OUT = WORK / "placeholder_audit.json"
PH = re.compile(r"\{\{([A-Z_]+)\}\}")


def collect(doc):
    out = [("body", p.text) for p in doc.paragraphs]

    def walk(t):
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    out.append(("table", p.text))
                for nt in cell.tables:
                    walk(nt)

    for t in doc.tables:
        walk(t)
    for s in doc.sections:
        out += [("header", p.text) for p in s.header.paragraphs]
        out += [("footer", p.text) for p in s.footer.paragraphs]
    return out


@pytest.fixture(scope="module")
def expected():
    doc = Document(str(TEMPLATE))
    data = json.loads(DATA.read_text())
    counts, locs, conds = {}, {}, set()
    for loc, text in collect(doc):
        for name in PH.findall(text):
            if name.startswith("IF_"):
                conds.add(name[3:])
                continue
            if name.startswith("END_IF_"):
                continue
            counts[name] = counts.get(name, 0) + 1
            locs.setdefault(name, set()).add(loc)
    return {
        "placeholders": {k: {"count": counts[k], "locations": sorted(locs[k])} for k in counts},
        "conditions": sorted(conds),
        "missing_keys": sorted(k for k in counts if k not in data),
        "unused_keys": sorted(k for k in data if k not in counts),
    }


@pytest.fixture(scope="module")
def got():
    assert OUT.exists(), f"missing {OUT}"
    return json.loads(OUT.read_text())


def test_schema(got):
    assert set(got.keys()) == {"placeholders", "conditions", "missing_keys", "unused_keys"}
    for k, v in got["placeholders"].items():
        assert set(v.keys()) == {"count", "locations"}, k
        assert set(v["locations"]) <= {"body", "table", "header", "footer"}, k


def test_placeholder_names(got, expected):
    assert sorted(got["placeholders"]) == sorted(expected["placeholders"])


def test_counts(got, expected):
    for k, v in expected["placeholders"].items():
        assert got["placeholders"][k]["count"] == v["count"], k


def test_locations(got, expected):
    for k, v in expected["placeholders"].items():
        assert list(got["placeholders"][k]["locations"]) == v["locations"], k


def test_conditions_and_keys(got, expected):
    assert list(got["conditions"]) == expected["conditions"]
    assert list(got["missing_keys"]) == expected["missing_keys"]
    assert list(got["unused_keys"]) == expected["unused_keys"]
