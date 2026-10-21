import json
import random

from docx import Document

random.seed(5309)
POOL = ["CANDIDATE_FULL_NAME", "POSITION", "DEPARTMENT", "START_DATE", "BASE_SALARY", "SIGNING_BONUS",
        "EQUITY_SHARES", "MANAGER_NAME", "HR_NAME", "PTO_DAYS", "CITY", "STATE", "ZIP_CODE",
        "RESPONSE_DEADLINE", "OFFER_ID", "COMPANY_NAME"]
FILL = ["Please", "note", "that", "your", "offer", "includes", "the", "following", "terms", "as", "agreed", "with"]
CONDS = ["RELOCATION", "SIGNING_BONUS", "REMOTE_WORK", "VISA_SPONSORSHIP"]
EXTRA = ["BENEFITS_TIER", "BADGE_ID", "OFFICE_FLOOR", "PARKING_SPOT"]

chosen = random.sample(POOL, 12)
locs = ["body", "table", "nested", "header", "footer"]
slots = {l: [] for l in locs}
slots["header"].append(chosen[0])
slots["footer"].append(chosen[1])
slots["nested"].append(chosen[2])
slots["table"].append(chosen[3])
slots["body"].extend(chosen[4:])
for name in chosen:
    for _ in range(random.randint(0, 2)):
        slots[random.choice(locs)].append(name)
for l in locs:
    random.shuffle(slots[l])


def add_split(p, text):
    k = random.randint(1, len(text) - 1)
    p.add_run(text[:k])
    p.add_run(text[k:])


def fill(p, names):
    for n in names:
        p.add_run(" ".join(random.sample(FILL, 3)) + " ")
        add_split(p, "{{" + n + "}}")
    p.add_run(".")


doc = Document()
sec = doc.sections[0]
fill(sec.header.paragraphs[0], slots["header"])
fill(sec.footer.paragraphs[0], slots["footer"])
doc.add_paragraph("Offer Letter")
body = slots["body"]
half = len(body) // 2
first, second = body[:half], body[half:]
for i in range(0, len(first), 2):
    fill(doc.add_paragraph(), first[i:i + 2])
for c in random.sample(CONDS, 2):
    add_split(doc.add_paragraph(), "{{IF_" + c + "}}")
    fill(doc.add_paragraph(), [random.choice(chosen)])
    add_split(doc.add_paragraph(), "{{END_IF_" + c + "}}")
for i in range(0, len(second), 3):
    fill(doc.add_paragraph(), second[i:i + 3])
t = doc.add_table(rows=3, cols=2)
top = [t.cell(0, 0), t.cell(0, 1), t.cell(1, 0), t.cell(1, 1), t.cell(2, 0)]
nt = t.cell(2, 1).add_table(rows=2, cols=2)
inner = [nt.cell(r, c) for r in range(2) for c in range(2)]
for i, n in enumerate(slots["table"]):
    fill(top[i % len(top)].add_paragraph(), [n])
for i, n in enumerate(slots["nested"]):
    fill(inner[i % len(inner)].add_paragraph(), [n])
doc.add_paragraph("Sincerely,")
fill(doc.add_paragraph(), [random.choice(chosen)])
doc.save("offer_letter_template.docx")

missing = random.sample(chosen, 2)
unused_pool = [n for n in POOL if n not in chosen]
keys = [n for n in chosen if n not in missing] + random.sample(EXTRA, 2) + random.sample(unused_pool, 1)
random.shuffle(keys)
data = {k: "value-%d" % random.randint(100, 999) for k in keys}
with open("employee_data.json", "w") as f:
    json.dump(data, f, indent=2)
