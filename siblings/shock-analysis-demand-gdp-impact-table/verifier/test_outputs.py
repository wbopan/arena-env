import json
import math
from pathlib import Path

import openpyxl
import pytest

W = Path("/root/work")
OUT = W / "shock_impact.xlsx"
COLS = ["Year", "Deflator", "Real_GDP", "Investment_GEL", "Real_Investment", "GDP_Impact", "Impact_pct"]


def expected():
    ws = openpyxl.load_workbook(W / "weo.xlsx", data_only=True)["WEO_Data"]
    d = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        d[int(r[0])] = r
    p = json.loads((W / "params.json").read_text())
    hist = sorted(y for y in d if d[y][1] is not None and d[y][3] is not None)
    L = hist[-1]
    g = sum(d[t][3] / d[t - 1][3] - 1 for t in range(L - 3, L + 1)) / 4
    D = d[L][3]
    R = d[L][1] / D * 100
    w = p["allocation_weights"]
    sw = sum(w)
    start = p["start_year"]
    rows = []
    for t in range(L + 1, start + len(w)):
        D *= 1 + g
        R *= 1 + d[t][2] / 100
        if t >= start:
            inv = p["total_investment_usd_bn"] * p["fx_gel_per_usd"] * w[t - start] / sw
            ri = inv / D * 100
            imp = ri * (1 - p["import_content_share"]) * p["demand_multiplier"]
            rows.append([t, D, R, inv, ri, imp, imp / R * 100])
    return rows


@pytest.fixture(scope="module")
def data():
    assert OUT.exists(), "output workbook missing"
    wb = openpyxl.load_workbook(OUT, data_only=True)
    assert "Impact" in wb.sheetnames, wb.sheetnames
    rows = [list(r) for r in wb["Impact"].iter_rows(values_only=True)]
    while rows and all(v is None for v in rows[-1]):
        rows.pop()
    return rows


def _close(a, b):
    return isinstance(a, (int, float)) and math.isclose(float(a), b, rel_tol=1e-5, abs_tol=1e-6)


def _check(data, idxs):
    exp = expected()
    body = data[1:]
    assert len(body) == len(exp)
    for got, e in zip(body, exp):
        for i in idxs:
            assert _close(got[i], e[i]), (e[0], COLS[i], got[i], e[i])


def test_headers(data):
    assert [str(v).strip() if v is not None else None for v in data[0][:7]] == COLS


def test_years(data):
    exp = expected()
    years = [r[0] for r in data[1:]]
    assert len(years) == len(exp)
    assert [int(y) for y in years] == [e[0] for e in exp]


def test_deflator_real_gdp(data):
    _check(data, [1, 2])


def test_investment(data):
    _check(data, [3, 4])


def test_impact(data):
    _check(data, [5, 6])
