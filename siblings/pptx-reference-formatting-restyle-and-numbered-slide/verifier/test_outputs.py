from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Pt
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn

IN = Path('/root/work/deck.pptx')
OUT = Path('/root/work/deck_processed.pptx')
TOL = 12700
INCH = 914400


def norm(s):
    return ' '.join(s.split())


def boxes(prs, limit, notes=False):
    res = []
    for sn, slide in enumerate(prs.slides, start=1):
        if sn > limit:
            break
        for sh in slide.shapes:
            if sh.is_placeholder or not sh.has_text_frame:
                continue
            t = norm(sh.text_frame.text)
            if not t:
                continue
            if t.startswith('Note:') == notes:
                res.append((sn, sh))
    return res


@pytest.fixture(scope='module')
def decks():
    assert OUT.exists(), 'output deck missing'
    src = Presentation(str(IN))
    out = Presentation(str(OUT))
    return src, out, len(src.slides)


def test_titles_preserved(decks):
    src, out, n = decks
    a = [(sn, norm(sh.text_frame.text)) for sn, sh in boxes(src, n)]
    b = [(sn, norm(sh.text_frame.text)) for sn, sh in boxes(out, n)]
    assert a and a == b


def test_notes_untouched(decks):
    src, out, n = decks
    key = lambda sn, sh: (sn, norm(sh.text_frame.text), sh.left, sh.top, sh.width, sh.height)
    a = [key(sn, sh) for sn, sh in boxes(src, n, notes=True)]
    b = [key(sn, sh) for sn, sh in boxes(out, n, notes=True)]
    assert a == b


def test_title_style(decks):
    src, out, n = decks
    found = boxes(out, n)
    assert found
    for sn, sh in found:
        for p in sh.text_frame.paragraphs:
            assert p.alignment == PP_ALIGN.CENTER, f'slide {sn}: not centered'
            for r in p.runs:
                f = r.font
                assert f.name == 'Arial', f'slide {sn}: font {f.name}'
                assert f.size is not None and abs(int(f.size) - int(Pt(16))) < 1, f'slide {sn}: size'
                assert f.bold in (None, False), f'slide {sn}: bold'
                try:
                    rgb = str(f.color.rgb)
                except AttributeError:
                    rgb = ''
                assert rgb.upper() == '989596', f'slide {sn}: color {rgb}'


def test_title_geometry(decks):
    src, out, n = decks
    W, H = out.slide_width, out.slide_height
    exp = (INCH // 2, H - 3 * INCH // 4, W - INCH, INCH // 2)
    for sn, sh in boxes(out, n):
        got = (sh.left, sh.top, sh.width, sh.height)
        for g, e in zip(got, exp):
            assert g is not None and abs(int(g) - int(e)) <= TOL, f'slide {sn}: {got} vs {exp}'


def _expected_refs(src, n):
    seen, exp = set(), []
    for sn, sh in boxes(src, n):
        t = norm(sh.text_frame.text)
        if t.lower() not in seen:
            seen.add(t.lower())
            exp.append(t)
    return exp


def _body_paras(out):
    last = out.slides[-1]
    body = [sh for sh in last.placeholders if sh.placeholder_format.idx == 1]
    assert body, 'no body placeholder on reference slide'
    return [p for p in body[0].text_frame.paragraphs if p.text.strip()]


def test_reference_content(decks):
    src, out, n = decks
    assert len(out.slides) == n + 1
    last = out.slides[-1]
    assert last.shapes.title is not None
    assert norm(last.shapes.title.text_frame.text) == 'Reference'
    assert [norm(p.text) for p in _body_paras(out)] == _expected_refs(src, n)


def test_reference_numbering(decks):
    src, out, n = decks
    paras = _body_paras(out)
    assert paras
    for p in paras:
        pPr = p._p.pPr
        assert pPr is not None, 'missing pPr'
        bu = pPr.find(qn('a:buAutoNum'))
        assert bu is not None and bu.get('type') == 'arabicPeriod'
