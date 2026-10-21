import random
import re
import zipfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

random.seed(4242)
PAPERS = [
    'Foam-Agent: Towards Automated Intelligent CFD Workflows',
    'ReAct: Synergizing Reasoning and Acting in Language Models',
    'Why Do Multi-Agent LLM Systems Fail?',
    'MultiAgentBench: Evaluating the Collaboration and Competition of LLM Agents',
    'Toolformer: Language Models Can Teach Themselves to Use Tools',
    'Reflexion: Language Agents with Verbal Reinforcement Learning',
    'Voyager: An Open-Ended Embodied Agent with Large Language Models',
    'CAMEL: Communicative Agents for Mind Exploration',
]
TOPICS = ['Planning', 'Tool Use', 'Memory', 'Benchmarks', 'Failure Modes', 'Embodiment']
picked = random.sample(PAPERS, 6)
seq = [picked[0], picked[1], picked[2], picked[1].lower(), picked[3], picked[4], picked[0].upper(), picked[5]]


def shrink(path):
    with zipfile.ZipFile(path) as z:
        items = [(n, z.read(n)) for n in z.namelist()]
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for n, d in items:
            if 'thumbnail' in n:
                continue
            if n == '_rels/.rels':
                d = re.sub(rb'<Relationship[^>]*thumbnail[^>]*/>', b'', d)
            z.writestr(n, d)


prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[0])
s.shapes.title.text = 'Awesome Agent Papers'
s.placeholders[1].text = 'Survey notes'
slots = sorted(random.sample(range(2, 12), 8))
notes = set(random.sample(range(2, 12), 3))
k = 0
for n in range(2, 12):
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = f'Section {n}: {random.choice(TOPICS)}'
    tf = s.placeholders[1].text_frame
    tf.text = f'Observation {random.randint(10, 99)}'
    tf.add_paragraph().text = f'Success rate {random.randint(20, 95)}%'
    if n in slots:
        b = s.shapes.add_textbox(Inches(random.uniform(0.2, 5.0)), Inches(random.uniform(1.0, 6.0)), Inches(random.uniform(1.5, 3.5)), Inches(random.uniform(0.6, 1.2)))
        b.text_frame.word_wrap = True
        r = b.text_frame.paragraphs[0].add_run()
        r.text = seq[k]
        k += 1
        r.font.name = random.choice(['Calibri', 'Georgia', 'Verdana'])
        r.font.size = Pt(random.choice([11, 12, 14, 18, 20]))
        r.font.bold = random.random() < 0.5
        r.font.color.rgb = RGBColor(random.randint(0, 255), random.randint(0, 255), random.randint(0, 255))
    if n in notes:
        b = s.shapes.add_textbox(Inches(6.5), Inches(0.2), Inches(3), Inches(0.4))
        r = b.text_frame.paragraphs[0].add_run()
        r.text = f'Note: draft figures v{random.randint(1, 9)}'
        r.font.size = Pt(10)
prs.save('deck.pptx')
shrink('deck.pptx')
