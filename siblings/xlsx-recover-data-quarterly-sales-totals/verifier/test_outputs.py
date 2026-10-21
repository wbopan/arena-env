import pytest
import openpyxl
from pathlib import Path

W = Path("/root/work")
IN = W / "sales_incomplete.xlsx"
OUT = W / "sales_recovered.xlsx"
SHEET = "Quarterly Sales"


def cons():
    c = [([(r, k) for k in range(2, 6)], (r, 6)) for r in range(2, 11)]
    c += [([(r, k) for r in range(2, 10)], (10, k)) for k in range(2, 7)]
    return c


def expected():
    ws = openpyxl.load_workbook(IN)[SHEET]
    vals = {}
    miss = set()
    for r in range(2, 11):
        for k in range(2, 7):
            v = ws.cell(row=r, column=k).value
            if v == "???":
                miss.add((r, k))
            else:
                vals[(r, k)] = v
    missing = sorted(miss)
    C = cons()
    while miss:
        prog = False
        for parts, tot in C:
            m = [x for x in parts + [tot] if x in miss]
            if len(m) == 1:
                x = m[0]
                s = sum(vals[p] for p in parts if p != x)
                vals[x] = s if x == tot else vals[tot] - s
                miss.discard(x)
                prog = True
        if not prog:
            raise RuntimeError("unsolvable input")
    return missing, vals


def isnum(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


@pytest.fixture(scope="module")
def out_ws():
    if not OUT.exists():
        pytest.fail("sales_recovered.xlsx not found")
    wb = openpyxl.load_workbook(OUT)
    assert SHEET in wb.sheetnames, f"sheet {SHEET} missing"
    return wb[SHEET]


def test_output_exists():
    assert OUT.exists()


def test_no_placeholders(out_ws):
    for row in out_ws.iter_rows(values_only=True):
        for v in row:
            assert v != "???", "placeholder remains"


def test_recovered_values(out_ws):
    missing, vals = expected()
    assert missing
    for (r, k) in missing:
        got = out_ws.cell(row=r, column=k).value
        assert isnum(got), f"cell r{r}c{k} not numeric: {got!r}"
        assert abs(got - vals[(r, k)]) < 1e-6, f"cell r{r}c{k}: expected {vals[(r, k)]}, got {got}"


def test_known_cells_unchanged(out_ws):
    ws_in = openpyxl.load_workbook(IN)[SHEET]
    for r in range(1, 11):
        for k in range(1, 7):
            v = ws_in.cell(row=r, column=k).value
            if v == "???":
                continue
            assert out_ws.cell(row=r, column=k).value == v, f"cell r{r}c{k} changed"


def test_constraints_hold(out_ws):
    g = lambda r, k: out_ws.cell(row=r, column=k).value
    for parts, tot in cons():
        assert all(isnum(g(*p)) for p in parts + [tot])
        assert abs(sum(g(*p) for p in parts) - g(*tot)) < 1e-6, f"constraint for {tot} fails"
