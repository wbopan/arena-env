import csv
import os
from collections import defaultdict

from openpyxl import load_workbook

DATA = "/root/work/data.xlsx"
OUT = "/root/work/commissions.csv"
TOL = 0.02


def rate_for(n):
    if n < 10000:
        return 0.003
    if n < 50000:
        return 0.002
    if n < 200000:
        return 0.0012
    return 0.0008


def expected():
    ws = load_workbook(DATA, data_only=True)["Trades"]
    tot = defaultdict(float)
    cnt = defaultdict(int)
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        client, side, q, p = r[1], str(r[2]).upper(), float(r[3]), float(r[4])
        n = round(q * p, 2)
        comm = max(n * rate_for(n), 7.5)
        fee = n * 0.000025 if side == "SELL" else 0.0
        tot[client] += round(comm + fee, 2)
        cnt[client] += 1
    return {c: (cnt[c], tot[c]) for c in tot}


def rows():
    with open(OUT, newline="") as f:
        return list(csv.reader(f))


def test_output_exists():
    assert os.path.exists(OUT)


def test_header():
    r = rows()
    assert [h.strip() for h in r[0]] == ["client", "trades", "total_charges"]


def test_clients_sorted_and_complete():
    r = [x for x in rows()[1:] if x]
    got = [x[0].strip() for x in r]
    assert got == sorted(expected().keys())


def test_trade_counts():
    e = expected()
    for x in rows()[1:]:
        if x:
            assert int(x[1]) == e[x[0].strip()][0], f"count mismatch for {x[0]}"


def test_total_charges():
    e = expected()
    for x in rows()[1:]:
        if x:
            v = float(x[2])
            exp = e[x[0].strip()][1]
            assert abs(v - exp) <= TOL, f"{x[0]}: expected {exp:.2f}, got {v}"
