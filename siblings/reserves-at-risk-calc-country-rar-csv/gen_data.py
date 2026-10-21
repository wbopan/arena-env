import csv
import math
import random
from openpyxl import Workbook

random.seed(23)
rows = []
p = 2050.0
for y in (2024, 2025):
    for m in range(1, 13):
        if y == 2025 and m > 9:
            break
        rows.append((f"{y}M{m:02d}", round(p, 2)))
        p *= math.exp(random.gauss(0.015, 0.04))
with open("gold_monthly.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(["period", "price_usd_per_ozt"])
    w.writerows(rows)

countries = ["Albania", "Brazil", "Chile", "Denmark", "Egypt", "Finland", "Ghana", "Hungary",
             "India", "Japan", "Kenya", "Latvia", "Mexico", "Norway", "Oman", "Peru"]
years = ["2023", "2024", "2025"]
wb = Workbook()

ws = wb.active
ws.title = "Value"
ws.append(["Country"] + years)
vc = countries[:10]
random.shuffle(vc)
for c in vc:
    base = random.uniform(500, 80000)
    vals = [round(base * random.uniform(0.85, 1.15), 2) for _ in years]
    if c in ("Denmark", "India"):
        vals[2] = None
    ws.append([c] + vals)

ws = wb.create_sheet("Volume")
ws.append(["Country"] + years)
vo = countries[2:14]
random.shuffle(vo)
for c in vo:
    base = random.uniform(0.5, 30)
    vals = [round(base * random.uniform(0.95, 1.05), 3) for _ in years]
    if c == "Latvia":
        vals[2] = None
    ws.append([c] + vals)

ws = wb.create_sheet("Total Reserves")
ws.append(["Country"] + years)
tr = [c for c in countries if c != "Kenya"]
random.shuffle(tr)
for c in tr:
    base = random.uniform(10000, 600000)
    vals = [round(base * random.uniform(0.9, 1.1), 1) for _ in years]
    if c == "Hungary":
        vals[2] = None
    ws.append([c] + vals)
wb.save("reserves.xlsx")
