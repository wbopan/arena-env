import csv
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

ROOT = Path('/root/work')
OUT = ROOT / 'spectra.xlsx'
HEADER = ['file_name', 'p_high_fraction', 's_high_fraction', 's_hv_ratio']


def reference():
    intervals = {}
    with (ROOT / 'windows.csv').open(newline='') as f:
        for r in csv.DictReader(f):
            intervals[(r['file_name'], r['phase'])] = slice(int(r['start_idx']), int(r['end_idx']))
    answer = {}
    for path in sorted(ROOT.glob('station_*.npz')):
        with np.load(path) as record:
            names = str(record['channels']).split(',')
            signal = record['data']
            dt = float(record['dt'])
        z = names.index(next(c for c in names if c.endswith('Z')))
        horizontal = [j for j, c in enumerate(names) if c.endswith(('N', 'E'))]
        p = signal[intervals[(path.name, 'P')]]
        s = signal[intervals[(path.name, 'S')]]
        p_power = np.abs(np.fft.rfft(p[:, z])) ** 2
        s_power = sum((np.abs(np.fft.rfft(s[:, j])) ** 2 for j in horizontal))
        pf = np.fft.rfftfreq(len(p), d=dt) >= 8
        sf = np.fft.rfftfreq(len(s), d=dt) >= 8
        ratio = np.sum(s[:, horizontal].astype('float64') ** 2) / np.sum(s[:, z].astype('float64') ** 2)
        answer[path.name] = [round(float(p_power[pf].sum() / p_power.sum()), 4), round(float(s_power[sf].sum() / s_power.sum()), 4), round(float(ratio), 4)]
    return answer


def table():
    book = load_workbook(OUT, read_only=True, data_only=True)
    try:
        return book.sheetnames, list(book['Spectra'].values)
    finally:
        book.close()


def test_file_exists():
    assert OUT.is_file()


def test_sheet_and_header():
    sheets, rows = table()
    assert sheets == ['Spectra']
    assert list(rows[0]) == HEADER


def test_station_rows():
    _, rows = table()
    assert [r[0] for r in rows[1:]] == sorted(reference())
    assert all(len(r) == 4 and all(type(v) in (int, float) for v in r[1:]) for r in rows[1:])


def test_spectral_fractions():
    _, rows = table()
    expected = reference()
    for r in rows[1:]:
        for actual, target in zip(r[1:3], expected[r[0]][:2]):
            assert abs(actual - target) <= 0.00005


def test_horizontal_vertical_ratio():
    _, rows = table()
    expected = reference()
    for r in rows[1:]:
        assert abs(r[3] - expected[r[0]][2]) <= 0.00005
