import json
from pathlib import Path
import numpy as np
import openpyxl
import pytest

W = Path('/root/work')
OUT = W / 'supply_trend.json'
TOL = 1e-5


def expected():
    wb = openpyxl.load_workbook(W / 'macro_data.xlsx', data_only=True)
    rows = [r for r in wb['Macro'].iter_rows(min_row=2, values_only=True) if r[0] is not None]
    rows.sort(key=lambda r: int(r[0]))
    prm = {str(r[0]).strip(): float(r[1]) for r in wb['Params'].iter_rows(min_row=2, values_only=True) if r[0] is not None}
    a = prm['alpha']
    lam = prm['hp_lambda']
    yrs = [int(r[0]) for r in rows]
    K = np.array([float(r[1]) for r in rows])
    Y = np.array([float(r[2]) for r in rows])
    C = np.array([float(r[3]) for r in rows])
    dep = float(np.mean((C / K)[-8:]))
    z = np.log(Y) - a * np.log(K)
    n = len(z)
    D = np.zeros((n - 2, n))
    for i in range(n - 2):
        D[i, i] = 1.0
        D[i, i + 1] = -2.0
        D[i, i + 2] = 1.0
    t = np.linalg.solve(np.eye(n) + lam * D.T @ D, z)
    return yrs, dep, z, t


@pytest.fixture(scope='module')
def out():
    assert OUT.exists(), 'supply_trend.json missing'
    return json.loads(OUT.read_text())


@pytest.fixture(scope='module')
def exp():
    return expected()


def test_file_exists():
    assert OUT.exists()


def test_schema(out, exp):
    yrs = exp[0]
    for k in ('depreciation_rate', 'trend_growth_last', 'series'):
        assert k in out, k
    s = out['series']
    assert len(s) == len(yrs)
    for e in s:
        assert {'year', 'lnZ', 'lnZ_trend'} <= set(e)
    assert [int(e['year']) for e in s] == yrs


def test_depreciation(out, exp):
    assert abs(float(out['depreciation_rate']) - exp[1]) < TOL


def test_lnz(out, exp):
    for e, v in zip(out['series'], exp[2]):
        assert abs(float(e['lnZ']) - v) < TOL, e['year']


def test_hp_trend(out, exp):
    for e, v in zip(out['series'], exp[3]):
        assert abs(float(e['lnZ_trend']) - v) < TOL, e['year']


def test_trend_growth(out, exp):
    t = exp[3]
    assert abs(float(out['trend_growth_last']) - (t[-1] - t[-2])) < TOL
