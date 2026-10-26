import random
from openpyxl import Workbook

random.seed(1842)
proteins = [f'PX{i:03d}' for i in range(1, 15)]
ids = [f'CASE{i:02d}' for i in range(1, 6)]
pairs = [(p, f'{p}_BREAST_Control', f'{p}_BREAST_Treated') for p in ids]
samples = [s for _, c, t in pairs for s in (c, t)]
random.shuffle(samples)
rows = {}
for j, protein in enumerate(proteins):
    baseline = random.uniform(7.0, 12.0)
    effect = random.uniform(-1.8, 1.8)
    values = {}
    for pair, control, treated in pairs:
        subject = random.uniform(-0.7, 0.7)
        values[control] = round(baseline + subject + random.uniform(-0.14, 0.14), 5)
        values[treated] = round(baseline + subject + effect + random.uniform(-0.35, 0.35), 5)
    rows[protein] = values
random.shuffle(proteins)
wb = Workbook()
data = wb.active
data.title = 'Data'
data.append(['Protein_ID', *samples])
for protein in proteins:
    data.append([protein, *(rows[protein][s] for s in samples)])
ws = wb.create_sheet('Pairs')
ws.append(['Pair_ID', 'Control_Sample', 'Treated_Sample'])
for pair in pairs:
    ws.append(pair)
ws = wb.create_sheet('Targets')
ws.append(['Protein_ID'])
for protein in random.sample(proteins, 7):
    ws.append([protein])
wb.save('paired_expression.xlsx')
