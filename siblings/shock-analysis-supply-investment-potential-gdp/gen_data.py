import random
from openpyxl import Workbook

random.seed(2026)
alpha = round(random.uniform(0.30, 0.42), 3)
delta = round(random.uniform(0.04, 0.08), 4)
wb = Workbook()
ws = wb.active
ws.title = 'Baseline'
ws.append(['year', 'K_base', 'lnZ_trend'])
k = random.uniform(90, 130)
z = random.uniform(1.6, 2.0)
g = random.uniform(0.008, 0.02)
for yr in range(2024, 2042):
    ws.append([yr, round(k, 4), round(z, 6)])
    k *= 1 + random.uniform(0.02, 0.05)
    z += g
iv = wb.create_sheet('Investment')
iv.append(['year', 'investment'])
rows = [[yr, round(random.uniform(0.4, 1.3), 4)] for yr in range(2026, 2034)]
random.shuffle(rows)
for r in rows:
    iv.append(r)
p = wb.create_sheet('Params')
p.append(['parameter', 'value'])
p.append(['alpha', alpha])
p.append(['delta', delta])
wb.save('shock_inputs.xlsx')
