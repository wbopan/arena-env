import math
from pathlib import Path
import openpyxl
import pytest

W = Path('/root/work')
OUT = W / 'shock_results.xlsx'
HDR = ['year', 'I', 'dK', 'K_with', 'Ystar_base', 'Ystar_with', 'gap_pct']


def expected():
    wb = openpyxl.load_workbook(W / 'shock_inputs.xlsx', data_only=True)
    b = sorted([r for r in wb['Baseline'].iter_rows(min_row=2, values_only=True) if r[0] is not None], key=lambda r: int(r[0]))
    inv = {}
    for r in wb['Investment'].iter_rows(min_row=2, values_only=True):
        if r[0] is not None:
            inv[int(r[0])] = inv.get(int(r[0]), 0.0) + float(r[1])
    prm = {str(r[0]).strip(): float(r[1]) for r in wb['Params'].iter_rows(min_row=2, values_only=True) if r[0] is not None}
    a, d = prm['alpha'], prm['delta']
    res = []
    dk = 0.0
    for yr, kb, z in b:
        kb = float(kb)
        z = float(z)
        i = inv.get(int(yr), 0.0)
        dk = (1 - d) * dk + i
        kw = kb + dk
        yb = math.exp(z) * kb ** a
        yw = math.exp(z) * kw ** a
        res.append([int(yr), i, dk, kw, yb, yw, 100 * (yw / yb - 1)])
    return res


@pytest.fixture(scope='module')
def out():
    assert OUT.exists(), 'shock_results.xlsx missing'
    wb = openpyxl.load_workbook(OUT, data_only=True)
    assert 'Shock' in wb.sheetnames
    rows = list(wb['Shock'].iter_rows(values_only=True))
    hdr = [str(h).strip() if h is not None else '' for h in rows[0]]
    data = [r for r in rows[1:] if r and r[0] is not None]
    return hdr, data


@pytest.fixture(scope='module')
def exp():
    return expected()


def close(a, b):
    return a is not None and math.isclose(float(a), b, rel_tol=1e-6, abs_tol=1e-5)


def test_file_and_sheet():
    assert OUT.exists()
    assert 'Shock' in openpyxl.load_workbook(OUT).sheetnames


def test_header(out):
    assert out[0][:7] == HDR


def test_years(out, exp):
    assert [int(r[0]) for r in out[1]] == [e[0] for e in exp]


def test_capital(out, exp):
    for r, e in zip(out[1], exp):
        for j in (1, 2, 3):
            assert close(r[j], e[j]), (e[0], HDR[j], r[j], e[j])


def test_potential_output(out, exp):
    for r, e in zip(out[1], exp):
        for j in (4, 5):
            assert close(r[j], e[j]), (e[0], HDR[j], r[j], e[j])


def test_gap(out, exp):
    for r, e in zip(out[1], exp):
        assert close(r[6], e[6]), (e[0], r[6], e[6])
