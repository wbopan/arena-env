import re
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

W = Path("/root/work")
OUT = W / "price_changes.xlsx"
TOL = 0.011


def old_prices():
    d = {}
    for page in PdfReader(str(W / "price_list_backup.pdf")).pages:
        for line in (page.extract_text() or "").splitlines():
            line = line.strip()
            if re.match(r"^SKU-\d{4}\s*\|", line):
                parts = [x.strip() for x in line.split("|")]
                d[parts[0]] = round(float(parts[2].replace(",", "")), 2)
    return d


def new_prices():
    ws = load_workbook(W / "price_list_current.xlsx", data_only=True)["Prices"]
    rows = list(ws.iter_rows(values_only=True))
    h = list(rows[0])
    i, j = h.index("SKU"), h.index("Unit Price")
    return {str(r[i]).strip(): round(float(r[j]), 2) for r in rows[1:] if r[i] is not None}


def expected():
    o, n = old_prices(), new_prices()
    res = []
    for s in sorted(set(o) | set(n)):
        if s not in n:
            res.append((s, "REMOVED", o[s], None, None))
        elif s not in o:
            res.append((s, "ADDED", None, n[s], None))
        elif o[s] != n[s]:
            res.append((s, "CHANGED", o[s], n[s], round((n[s] - o[s]) / o[s] * 100, 2)))
    return res


def read_output():
    wb = load_workbook(OUT, data_only=True)
    assert "Changes" in wb.sheetnames, f"Sheet 'Changes' missing, found {wb.sheetnames}"
    rows = [r for r in wb["Changes"].iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
    return rows


def blank(v):
    return v is None or (isinstance(v, str) and v.strip() == "")


def close(a, b):
    if b is None:
        return blank(a)
    try:
        return abs(float(a) - b) <= TOL
    except (TypeError, ValueError):
        return False


def test_output_exists():
    assert OUT.exists(), f"{OUT} not found"


def test_sheet_and_header():
    rows = read_output()
    assert rows, "Changes sheet is empty"
    hdr = [str(v).strip() if v is not None else "" for v in rows[0][:5]]
    assert hdr == ["SKU", "Status", "Old Price", "New Price", "Change Pct"], hdr


def test_skus_status_and_order():
    rows = read_output()[1:]
    got = [(str(r[0]).strip(), str(r[1]).strip().upper()) for r in rows]
    exp = [(e[0], e[1]) for e in expected()]
    assert got == exp, f"Expected {exp}, got {got}"


def test_price_values():
    rows = {str(r[0]).strip(): r for r in read_output()[1:]}
    for s, st, op, npr, _ in expected():
        assert s in rows, f"{s} missing"
        r = rows[s]
        assert close(r[2], op), f"{s}: Old Price {r[2]!r}, expected {op}"
        assert close(r[3], npr), f"{s}: New Price {r[3]!r}, expected {npr}"


def test_change_pct():
    rows = {str(r[0]).strip(): r for r in read_output()[1:]}
    for s, st, _, _, pct in expected():
        assert s in rows, f"{s} missing"
        v = rows[s][4] if len(rows[s]) > 4 else None
        assert close(v, pct), f"{s}: Change Pct {v!r}, expected {pct}"
