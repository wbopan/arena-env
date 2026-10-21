import json
import pytest
import openpyxl
from pathlib import Path

W = Path("/root/work")
IN = W / "budget_incomplete.xlsx"
OUT = W / "recovered.json"


def comp(c, V, m):
    t, i, k = c
    kn = lambda x: x not in m
    if t == "B":
        v = None
        if k > 0 and kn(("B", i, k - 1)) and kn(("Y", i, k)):
            v = V[("B", i, k - 1)] * (1 + V[("Y", i, k)] / 100)
        elif k < 5 and kn(("B", i, k + 1)) and kn(("Y", i, k + 1)):
            v = V[("B", i, k + 1)] / (1 + V[("Y", i, k + 1)] / 100)
        elif k == 0 and kn(("B", i, 5)) and kn(("C", i, 0)):
            v = V[("B", i, 5)] / (1 + V[("C", i, 0)] / 100) ** 5
        elif k == 5 and kn(("B", i, 0)) and kn(("C", i, 0)):
            v = V[("B", i, 0)] * (1 + V[("C", i, 0)] / 100) ** 5
        elif kn(("A", i, 0)) and all(kn(("B", i, j)) for j in range(6) if j != k):
            v = 6 * V[("A", i, 0)] - sum(V[("B", i, j)] for j in range(6) if j != k)
        return None if v is None else int(round(v))
    if t == "Y":
        if kn(("B", i, k)) and kn(("B", i, k - 1)):
            return (V[("B", i, k)] / V[("B", i, k - 1)] - 1) * 100
        return None
    if t == "C":
        if kn(("B", i, 0)) and kn(("B", i, 5)):
            return ((V[("B", i, 5)] / V[("B", i, 0)]) ** 0.2 - 1) * 100
        return None
    if all(kn(("B", i, j)) for j in range(6)):
        return sum(V[("B", i, j)] for j in range(6)) / 6
    return None


def key(c):
    t, i, k = c
    if t == "B":
        return f"Budget!{'BCDEFG'[k]}{i + 2}"
    if t == "Y":
        return f"YoY (%)!{'BCDEF'[k - 1]}{i + 2}"
    if t == "C":
        return f"Summary!B{i + 2}"
    return f"Summary!C{i + 2}"


def expected():
    wb = openpyxl.load_workbook(IN)
    b, y, s = wb["Budget"], wb["YoY (%)"], wb["Summary"]
    N = 0
    while b.cell(row=N + 2, column=1).value not in (None, ""):
        N += 1
    V = {}
    for i in range(N):
        for k in range(6):
            V[("B", i, k)] = b.cell(row=i + 2, column=k + 2).value
        for k in range(1, 6):
            V[("Y", i, k)] = y.cell(row=i + 2, column=k + 1).value
        V[("C", i, 0)] = s.cell(row=i + 2, column=2).value
        V[("A", i, 0)] = s.cell(row=i + 2, column=3).value
    miss = {c for c, v in V.items() if v == "???"}
    missing = sorted(miss)
    while miss:
        prog = False
        for c in sorted(miss):
            v = comp(c, V, miss)
            if v is not None:
                V[c] = v
                miss.discard(c)
                prog = True
        if not prog:
            raise RuntimeError("unsolvable input")
    return {key(c): (c[0], V[c]) for c in missing}


@pytest.fixture(scope="module")
def data():
    if not OUT.exists():
        pytest.fail("recovered.json not found")
    d = json.loads(OUT.read_text())
    assert isinstance(d, dict)
    return d


@pytest.fixture(scope="module")
def exp():
    return expected()


def isnum(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def test_output_is_json_object(data):
    assert len(data) > 0
    assert all(isnum(v) for v in data.values()), "all values must be numbers"


def test_keys_match_missing_cells(data, exp):
    assert set(data.keys()) == set(exp.keys()), f"expected keys {sorted(exp)}, got {sorted(data)}"


def test_budget_values(data, exp):
    for k, (t, v) in exp.items():
        if t == "B":
            assert k in data and isnum(data[k]), f"{k} missing"
            assert abs(data[k] - v) <= 0.5, f"{k}: expected {v}, got {data[k]}"


def test_percentage_values(data, exp):
    for k, (t, v) in exp.items():
        if t in ("Y", "C"):
            assert k in data and isnum(data[k]), f"{k} missing"
            assert abs(data[k] - round(v, 2)) <= 0.011, f"{k}: expected {round(v, 2)}, got {data[k]}"


def test_average_values(data, exp):
    for k, (t, v) in exp.items():
        if t == "A":
            assert k in data and isnum(data[k]), f"{k} missing"
            assert abs(data[k] - round(v, 2)) <= 0.011, f"{k}: expected {round(v, 2)}, got {data[k]}"
