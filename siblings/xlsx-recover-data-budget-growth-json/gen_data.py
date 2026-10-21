import random
from openpyxl import Workbook

random.seed(9031)
depts = ["Science", "Exploration", "Space Operations", "Technology", "Aeronautics"]
N = len(depts)
B = []
for d in depts:
    v = random.randint(800, 8000)
    row = [v]
    for k in range(5):
        v = int(round(v * (1 + random.uniform(-0.06, 0.12))))
        row.append(v)
    B.append(row)
Y = {(i, k): round((B[i][k] / B[i][k - 1] - 1) * 100, 4) for i in range(N) for k in range(1, 6)}
C = {i: round(((B[i][5] / B[i][0]) ** 0.2 - 1) * 100, 4) for i in range(N)}
A = {i: round(sum(B[i]) / 6, 2) for i in range(N)}

cells = [("B", i, k) for i in range(N) for k in range(6)]
cells += [("Y", i, k) for i in range(N) for k in range(1, 6)]
cells += [("C", i, 0) for i in range(N)] + [("A", i, 0) for i in range(N)]


def can(c, m):
    t, i, k = c
    kn = lambda x: x not in m
    if t == "B":
        if k > 0 and kn(("B", i, k - 1)) and kn(("Y", i, k)):
            return True
        if k < 5 and kn(("B", i, k + 1)) and kn(("Y", i, k + 1)):
            return True
        if k == 0 and kn(("B", i, 5)) and kn(("C", i, 0)):
            return True
        if k == 5 and kn(("B", i, 0)) and kn(("C", i, 0)):
            return True
        return kn(("A", i, 0)) and all(kn(("B", i, j)) for j in range(6) if j != k)
    if t == "Y":
        return kn(("B", i, k)) and kn(("B", i, k - 1))
    js = (0, 5) if t == "C" else range(6)
    return all(kn(("B", i, j)) for j in js)


def rounds(miss):
    m = set(miss)
    r = 0
    while m:
        new = {c for c in m if can(c, m)}
        if not new:
            return None
        m -= new
        r += 1
    return r


while True:
    M = random.sample(cells, 10)
    nb = sum(c[0] == "B" for c in M)
    ny = sum(c[0] == "Y" for c in M)
    ns = sum(c[0] in "CA" for c in M)
    if nb >= 4 and ny >= 2 and ns >= 2:
        r = rounds(M)
        if r is not None and r >= 2:
            break
M = set(M)

wb = Workbook()
ws = wb.active
ws.title = "Budget"
ws.append(["Department"] + [f"FY{y}" for y in range(2019, 2025)])
for i, d in enumerate(depts):
    ws.append([d] + [("???" if ("B", i, k) in M else B[i][k]) for k in range(6)])
wy = wb.create_sheet("YoY (%)")
wy.append(["Department"] + [f"FY{y}" for y in range(2020, 2025)])
for i, d in enumerate(depts):
    wy.append([d] + [("???" if ("Y", i, k) in M else Y[(i, k)]) for k in range(1, 6)])
wsum = wb.create_sheet("Summary")
wsum.append(["Department", "CAGR FY2019-FY2024 (%)", "Average Budget FY2019-FY2024"])
for i, d in enumerate(depts):
    wsum.append([d, "???" if ("C", i, 0) in M else C[i], "???" if ("A", i, 0) in M else A[i]])
wb.save("budget_incomplete.xlsx")
