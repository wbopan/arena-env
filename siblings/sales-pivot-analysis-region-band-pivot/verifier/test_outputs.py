import os
from collections import defaultdict

import pytest
from openpyxl import load_workbook

W = "/root/work"
OUT = os.path.join(W, "region_report.xlsx")
HEADER = ["OrderID", "Region", "Rep", "Units", "UnitPrice", "Revenue", "SizeBand"]


def band(u):
    if u < 10:
        return "Small"
    if u <= 49:
        return "Medium"
    return "Large"


@pytest.fixture(scope="module")
def source():
    ws = load_workbook(os.path.join(W, "sales.xlsx"))["Orders"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = list(rows[0])
    out = []
    for r in rows[1:]:
        if r[0] is None:
            continue
        d = dict(zip(hdr, r))
        d["Revenue"] = round(d["Units"] * d["UnitPrice"], 2)
        d["SizeBand"] = band(d["Units"])
        out.append(d)
    return out


@pytest.fixture(scope="module")
def wb():
    assert os.path.exists(OUT), "region_report.xlsx missing"
    return load_workbook(OUT)


def test_sheets_exist(wb):
    for s in ["SourceData", "Revenue by Region", "Region Totals"]:
        assert s in wb.sheetnames, f"missing sheet {s}"


def test_sourcedata_schema(wb, source):
    rows = [r for r in wb["SourceData"].iter_rows(values_only=True) if any(v is not None for v in r)]
    assert [str(h).strip() for h in rows[0][:7]] == HEADER
    assert len(rows) - 1 == len(source)


def test_sourcedata_values(wb, source):
    rows = [r for r in wb["SourceData"].iter_rows(min_row=2, values_only=True) if r[0] is not None]
    for got, exp in zip(rows, source):
        assert str(got[0]) == exp["OrderID"]
        assert str(got[1]) == exp["Region"]
        assert abs(float(got[5]) - exp["Revenue"]) <= 0.011, f"Revenue wrong for {exp['OrderID']}"
        assert str(got[6]).strip() == exp["SizeBand"], f"SizeBand wrong for {exp['OrderID']}"


def test_pivot_structure(wb):
    ws = wb["Revenue by Region"]
    assert len(ws._pivots) >= 1, "no pivot table"
    p = ws._pivots[0]
    names = [f.name for f in p.cache.cacheFields]
    assert names[p.rowFields[0].x] == "Region"
    assert len(p.colFields) >= 1 and names[p.colFields[0].x] == "SizeBand"
    df = p.dataFields[0]
    assert names[df.fld] == "Revenue"
    assert (df.subtotal or "sum") == "sum"


def test_pivot_source_range(wb, source):
    p = wb["Revenue by Region"]._pivots[0]
    src = p.cache.cacheSource.worksheetSource
    assert src.sheet == "SourceData"
    ref = src.ref.replace("$", "").upper()
    start, end = ref.split(":")
    assert start == "A1"
    assert end == f"G{len(source) + 1}"


def test_region_totals(wb, source):
    tot = defaultdict(float)
    for d in source:
        tot[d["Region"]] += d["Revenue"]
    exp = sorted(((k, round(v, 2)) for k, v in tot.items()), key=lambda x: -x[1])
    rows = [r for r in wb["Region Totals"].iter_rows(values_only=True) if r and r[0] is not None]
    assert [str(h).strip() for h in rows[0][:2]] == ["Region", "Revenue"]
    body = rows[1:]
    assert len(body) == len(exp)
    for got, (reg, val) in zip(body, exp):
        assert str(got[0]) == reg
        assert abs(float(got[1]) - val) <= 0.02
