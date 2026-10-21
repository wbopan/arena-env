import random
from openpyxl import Workbook

random.seed(4127)
NAMES = ["Agriculture products", "Mining products", "Food products", "Textiles", "Wood products", "Chemicals", "Pharmaceuticals", "Basic metals", "Fabricated metal products", "Computers and electronics", "Electrical equipment", "Machinery and equipment", "Motor vehicles", "Other transport equipment", "Furniture", "Electricity and gas", "Construction work", "Trade services", "IT services", "Professional services"]
CAPITAL = {"Wood products", "Fabricated metal products", "Computers and electronics", "Electrical equipment", "Machinery and equipment", "Motor vehicles", "Other transport equipment", "Furniture", "Construction work", "IT services", "Professional services"}
SH = ["Code", "Product", "Output at basic prices", "Imports CIF", "Trade and transport margins", "Taxes less subsidies on products", "Total supply at purchasers prices"]
UH = ["Code", "Product", "Intermediate consumption", "Household final consumption", "Government final consumption", "Gross fixed capital formation", "Changes in inventories", "Exports", "Total use at purchasers prices"]

wb = Workbook()
s = wb.active
s.title = "SUPPLY"
u = wb.create_sheet("USE")
s.cell(1, 1, "Supply table at current prices, million GEL")
u.cell(1, 1, "Use table at current prices, million GEL")
for j, h in enumerate(SH, 1):
    s.cell(3, j, h)
for j, h in enumerate(UH, 1):
    u.cell(3, j, h)
stot = [0.0] * 5
utot = [0.0] * 7
r = 4
for i, name in enumerate(NAMES):
    code = f"P{i + 1:02d}"
    ic = round(random.uniform(200, 3000), 1)
    hh = round(random.uniform(0, 2500), 1)
    gov = round(random.uniform(0, 600), 1) if random.random() < 0.4 else 0.0
    if name in CAPITAL:
        gf = round(random.uniform(150, 2500), 1)
    else:
        gf = round(random.uniform(0, 120), 1) if random.random() < 0.3 else 0.0
    ch = round(random.uniform(-40, 80), 1)
    ex = round(random.uniform(0, 1500), 1)
    uses = [ic, hh, gov, gf, ch, ex]
    tot = round(sum(uses), 1)
    imp = round(tot * random.uniform(0.05, 0.75), 1)
    mg = round(tot * random.uniform(0, 0.08), 1)
    tx = round(tot * random.uniform(0, 0.06), 1)
    out = round(tot - imp - mg - tx, 1)
    sup = [out, imp, mg, tx, tot]
    s.cell(r, 1, code)
    s.cell(r, 2, name)
    u.cell(r, 1, code)
    u.cell(r, 2, name)
    for j, v in enumerate(sup):
        s.cell(r, 3 + j, v)
        stot[j] += v
    for j, v in enumerate(uses + [tot]):
        u.cell(r, 3 + j, v)
        utot[j] += v
    r += 1
s.cell(r, 1, "TOTAL")
s.cell(r, 2, "Total")
u.cell(r, 1, "TOTAL")
u.cell(r, 2, "Total")
for j, v in enumerate(stot):
    s.cell(r, 3 + j, round(v, 1))
for j, v in enumerate(utot):
    u.cell(r, 3 + j, round(v, 1))
wb.save("sut.xlsx")
