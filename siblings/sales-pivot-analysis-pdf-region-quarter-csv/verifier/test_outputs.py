import csv
import os
import re
from collections import defaultdict

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

W = "/root/work"
OUT = os.path.join(W, "region_quarter.csv")
PAT = re.compile(r"(\d{4})-(\d{2})-(\d{2})\s+(S\d{2})\s+([A-Za-z]+)\s+(-?\d+\.\d{2})")


@pytest.fixture(scope="module")
def expected():
    ws = load_workbook(os.path.join(W, "stores.xlsx"))["Stores"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(h) for h in rows[0]]
    si, ri = hdr.index("StoreID"), hdr.index("Region")
    smap = {r[si]: r[ri] for r in rows[1:] if r[si]}
    text = "\n".join(pg.extract_text() or "" for pg in PdfReader(os.path.join(W, "transactions.pdf")).pages)
    tx = PAT.findall(text)
    assert len(tx) > 50
    m = {reg: [0.0] * 4 for reg in set(smap.values())}
    for y, mo, d, st, cat, amt in tx:
        m[smap[st]][(int(mo) - 1) // 3] += float(amt)
    out = {}
    for reg in sorted(m):
        q = [round(v, 2) for v in m[reg]]
        out[reg] = q + [round(sum(m[reg]), 2)]
    gt = [round(sum(out[r][i] for r in out), 2) for i in range(5)]
    return out, gt


@pytest.fixture(scope="module")
def rows():
    assert os.path.exists(OUT), "region_quarter.csv missing"
    with open(OUT, newline="") as f:
        return [r for r in csv.reader(f) if r]


def test_header(rows):
    assert [c.strip() for c in rows[0]] == ["Region", "Q1", "Q2", "Q3", "Q4", "Total"]


def test_row_labels(rows, expected):
    out, _ = expected
    labels = [r[0].strip() for r in rows[1:]]
    assert labels == sorted(out) + ["Grand Total"]


def test_number_format(rows):
    for r in rows[1:]:
        for v in r[1:]:
            assert re.fullmatch(r"-?\d+\.\d{2}", v.strip()), f"bad number format {v!r}"


def test_region_values(rows, expected):
    out, _ = expected
    for r in rows[1:-1]:
        exp = out[r[0].strip()]
        for g, e in zip(r[1:6], exp):
            assert abs(float(g) - e) <= 0.02, f"{r[0]}: got {g}, expected {e}"


def test_grand_total(rows, expected):
    _, gt = expected
    last = rows[-1]
    assert last[0].strip() == "Grand Total"
    for g, e in zip(last[1:6], gt):
        assert abs(float(g) - e) <= 0.05, f"Grand Total: got {g}, expected {e}"
