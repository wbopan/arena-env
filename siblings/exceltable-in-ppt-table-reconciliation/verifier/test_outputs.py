import csv
import io
import os
import re
import zipfile

import pytest
from openpyxl import load_workbook
from pptx import Presentation

PPTX = '/root/work/budget.pptx'
OUT = '/root/work/mismatches.csv'
HEADER = ['department', 'column', 'slide_value', 'workbook_value']


def _num(ws, v):
    if isinstance(v, str) and v.startswith('='):
        m = re.fullmatch(r'=SUM[(]([A-Z]+[0-9]+:[A-Z]+[0-9]+)[)]', v.replace(' ', '').replace('$', '').upper())
        assert m, v
        return sum(_num(ws, c.value) for row in ws[m.group(1)] for c in row)
    return int(round(float(v)))


@pytest.fixture(scope='module')
def expected():
    with zipfile.ZipFile(PPTX) as z:
        names = [n for n in z.namelist() if n.startswith('ppt/embeddings/') and n.lower().endswith('.xlsx')]
        ws = load_workbook(io.BytesIO(z.read(names[0]))).worksheets[0]
    hdr = [str(c.value).strip() for c in ws[1]]
    book = {}
    for r in range(2, ws.max_row + 1):
        d = ws.cell(r, 1).value
        if d is None:
            continue
        book[str(d).strip()] = {hdr[c - 1]: _num(ws, ws.cell(r, c).value) for c in range(2, len(hdr) + 1)}
    table = None
    for sh in Presentation(PPTX).slides[0].shapes:
        if sh.has_table:
            table = sh.table
    assert table is not None
    th = [table.cell(0, j).text.strip() for j in range(len(table.columns))]
    rows = []
    for i in range(1, len(table.rows)):
        d = table.cell(i, 0).text.strip()
        for j in range(1, len(th)):
            sv = int(table.cell(i, j).text.replace(',', '').strip())
            wv = book[d][th[j]]
            if sv != wv:
                rows.append([d, th[j], str(sv), str(wv)])
    return rows


@pytest.fixture(scope='module')
def out():
    assert os.path.exists(OUT), 'mismatches.csv missing'
    with open(OUT, newline='') as f:
        return [[x.strip() for x in r] for r in csv.reader(f) if r]


def test_file_exists():
    assert os.path.exists(OUT)


def test_header(out):
    assert out[0] == HEADER


def test_row_count(out, expected):
    assert len(out) - 1 == len(expected), f'expected {len(expected)} mismatches, got {len(out) - 1}'


def test_mismatch_set(out, expected):
    norm = lambda r: (r[0], r[1], int(float(r[2])), int(float(r[3])))
    assert set(map(norm, out[1:])) == set(map(norm, expected))


def test_order(out, expected):
    assert [(r[0], r[1]) for r in out[1:]] == [(r[0], r[1]) for r in expected]
