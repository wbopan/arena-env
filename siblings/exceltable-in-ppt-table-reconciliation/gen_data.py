import io
import random
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import PROG_ID
from PIL import Image

random.seed(11)
depts = ['Engineering', 'Marketing', 'Sales', 'Finance', 'Operations', 'Legal', 'Support']
data = {d: [random.randint(8, 95) * 1000 + random.randint(0, 999) for _ in range(4)] for d in depts}
cols = ['Department', 'Q1', 'Q2', 'Q3', 'Q4', 'Total']

wb = Workbook()
ws = wb.active
ws.title = 'Budget'
ws.append(cols)
order = depts[:]
random.shuffle(order)
for k, d in enumerate(order):
    r = k + 2
    ws.append([d] + data[d] + [f'=SUM(B{r}:E{r})'])
xb = io.BytesIO()
wb.save(xb)
xb.seek(0)

icon = io.BytesIO()
Image.new('RGB', (32, 32), (33, 115, 70)).save(icon, 'PNG')
icon.seek(0)

vals = {d: data[d] + [sum(data[d])] for d in depts}
cells = random.sample([(i, j) for i in range(len(depts)) for j in range(5)], 5)
for i, j in cells:
    vals[depts[i]][j] += random.choice([-1, 1]) * random.randint(100, 4000)

prs = Presentation()
m = prs.slide_masters[0]
for l in list(m.slide_layouts):
    if l.name != 'Blank':
        m.slide_layouts.remove(l)
slide = prs.slides.add_slide(m.slide_layouts[0])
tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(9), Inches(0.7))
tb.text_frame.text = 'FY Budget by Department'
tb.text_frame.paragraphs[0].runs[0].font.size = Pt(26)
tbl = slide.shapes.add_table(len(depts) + 1, 6, Inches(0.4), Inches(1.0), Inches(7.2), Inches(3.6)).table
for j, c in enumerate(cols):
    tbl.cell(0, j).text = c
for i, d in enumerate(depts):
    tbl.cell(i + 1, 0).text = d
    for j in range(5):
        tbl.cell(i + 1, j + 1).text = f'{vals[d][j]:,}'
slide.shapes.add_ole_object(xb, PROG_ID.XLSX, Inches(8.0), Inches(1.0), Inches(1.5), Inches(1.5), icon_file=icon, icon_width=Inches(1), icon_height=Inches(1))
prs.save('budget.pptx')
