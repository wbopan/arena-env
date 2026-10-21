import io
import json
import os
import re
import zipfile

import pytest
from openpyxl import load_workbook
from pptx import Presentation

PPTX = '/root/work/report.pptx'
OUT = '/root/work/converted.json'


def _sheet():
    with zipfile.ZipFile(PPTX) as z:
        names = [n for n in z.namelist() if n.startswith('ppt/embeddings/') and n.lower().endswith('.xlsx')]
        assert names, 'no embedded workbook in input'
        return load_workbook(io.BytesIO(z.read(names[0])), data_only=False).worksheets[0]


def _val(ws, coord, depth=0):
    v = ws[coord].value
    if isinstance(v, str) and v.startswith('='):
        m = re.fullmatch(r'=1/[$]?([A-Z]+)[$]?([0-9]+)', v.replace(' ', ''))
        assert m and depth < 5
        return 1.0 / _val(ws, m.group(1) + m.group(2), depth + 1)
    return float(v)


@pytest.fixture(scope='module')
def expected():
    ws = _sheet()
    cols = {str(ws.cell(1, c).value).strip(): c for c in range(2, ws.max_column + 1) if ws.cell(1, c).value}
    rows = {str(ws.cell(r, 1).value).strip(): r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
    text = None
    for sh in Presentation(PPTX).slides[0].shapes:
        if sh.has_text_frame and sh.text_frame.text.strip().startswith('Report currency:'):
            text = sh.text_frame.text
    assert text
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    tgt = lines[0].split(':', 1)[1].strip()
    conv = {}
    for l in lines[1:]:
        name, rest = l.split(':', 1)
        amt_s, code = rest.split()
        amt = float(amt_s.replace(',', ''))
        rate = 1.0 if code == tgt else _val(ws, ws.cell(rows[code], cols[tgt]).coordinate)
        conv[name.strip()] = round(amt * rate, 2)
    return tgt, conv, round(sum(conv.values()), 2)


@pytest.fixture(scope='module')
def out():
    assert os.path.exists(OUT), 'converted.json missing'
    with open(OUT) as f:
        return json.load(f)


def test_file_exists_and_schema(out):
    assert isinstance(out, dict)
    for k in ('report_currency', 'converted', 'total'):
        assert k in out, k
    assert isinstance(out['converted'], dict)


def test_report_currency(out, expected):
    assert str(out['report_currency']).strip() == expected[0]


def test_region_names(out, expected):
    assert set(out['converted'].keys()) == set(expected[1].keys())


def test_converted_values(out, expected):
    for k, v in expected[1].items():
        got = out['converted'][k]
        assert isinstance(got, (int, float)), k
        assert abs(got - v) <= 0.011, f'{k}: expected {v}, got {got}'


def test_total(out, expected):
    assert isinstance(out['total'], (int, float))
    assert abs(out['total'] - expected[2]) <= 0.02
