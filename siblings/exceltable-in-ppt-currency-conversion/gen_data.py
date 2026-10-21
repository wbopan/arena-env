import io
import random
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import PROG_ID
from PIL import Image

random.seed(7)
cur = ['USD', 'EUR', 'GBP', 'JPY', 'CHF', 'CAD']
base = {'USD': 1.0, 'EUR': 1.08, 'GBP': 1.27, 'JPY': 0.0067, 'CHF': 1.12, 'CAD': 0.73}
u = {c: (v if c == 'USD' else v * random.uniform(0.95, 1.05)) for c, v in base.items()}
t = random.choice([1, 2, 3])

wb = Workbook()
ws = wb.active
ws.title = 'Rates'
ws.cell(1, 1, 'FROM/TO')
for k, c in enumerate(cur):
    ws.cell(1, k + 2, c)
    ws.cell(k + 2, 1, c)
for i, r in enumerate(cur):
    for j, c in enumerate(cur):
        if i == j:
            v = 1
        elif i > j and (j == t or random.random() < 0.5):
            v = '=1/' + get_column_letter(i + 2) + str(j + 2)
        else:
            v = round(u[r] / u[c], 6)
        ws.cell(i + 2, j + 2, v)
xb = io.BytesIO()
wb.save(xb)
xb.seek(0)

icon = io.BytesIO()
Image.new('RGB', (32, 32), (33, 115, 70)).save(icon, 'PNG')
icon.seek(0)

prs = Presentation()
m = prs.slide_masters[0]
for l in list(m.slide_layouts):
    if l.name != 'Blank':
        m.slide_layouts.remove(l)
slide = prs.slides.add_slide(m.slide_layouts[0])
tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.8))
tb.text_frame.text = 'Quarterly Regional Revenue'
tb.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
slide.shapes.add_ole_object(xb, PROG_ID.XLSX, Inches(0.5), Inches(1.5), Inches(4), Inches(3), icon_file=icon, icon_width=Inches(1), icon_height=Inches(1))

regions = ['Northwind', 'Southgate', 'Eastfield', 'Westbrook', 'Midland', 'Harbor']
cs = cur[:]
random.shuffle(cs)
lines = ['Report currency: ' + cur[t]]
for name, c in zip(regions, cs):
    amt = round(random.uniform(20000, 900000), 2)
    if c == 'JPY':
        amt = round(amt * 150, 2)
    lines.append(f'{name}: {amt:,.2f} {c}')
box = slide.shapes.add_textbox(Inches(5), Inches(1.5), Inches(4.5), Inches(3.5))
tf = box.text_frame
tf.text = lines[0]
for ln in lines[1:]:
    tf.add_paragraph().text = ln
prs.save('report.pptx')
