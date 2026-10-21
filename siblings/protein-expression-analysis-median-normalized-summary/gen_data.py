import random
import numpy as np
from openpyxl import Workbook

random.seed(11)
rng = np.random.default_rng(11)
lines = [("MDAMB468", "BREAST"), ("HCT116", "LARGE_INTESTINE"), ("A549", "LUNG"), ("MCF7", "BREAST"), ("PC3", "PROSTATE"), ("SKOV3", "OVARY")]
short = [f"TenPx{i+1:02d}" for i in range(12)]
full = [f"{lines[i % 6][0]}_{lines[i % 6][1]}_{s}" for i, s in enumerate(short)]
cond = ["Control"] * 6 + ["Treated"] * 6
random.shuffle(cond)
prot = ["Q%05d" % x for x in random.sample(range(10000, 99999), 40)]
base = rng.normal(18, 2, 40)
offs = rng.normal(0, 0.8, 12)
lg = base[:, None] + offs[None, :] + rng.normal(0, 0.35, (40, 12))
eff = np.zeros(40)
idx = rng.choice(40, 12, replace=False)
eff[idx] = rng.uniform(-2.5, 2.5, 12)
tr = np.array([c == "Treated" for c in cond])
lg[:, tr] += eff[:, None]
raw = np.round(2 ** lg).astype(int)
miss = rng.random((40, 12)) < 0.07
blank = rng.random((40, 12)) < 0.5
targets = random.sample(range(40), 8)
for ti in targets:
    for grp in (tr, ~tr):
        cols = np.where(grp)[0]
        while miss[ti, cols].sum() > 2:
            j = cols[miss[ti, cols]][0]
            miss[ti, j] = False

wb = Workbook()
ws = wb.active
ws.title = "Data"
perm = list(range(12))
random.shuffle(perm)
ws.append(["Protein_ID"] + [full[j] for j in perm])
for i in range(40):
    row = [prot[i]]
    for j in perm:
        if miss[i, j]:
            row.append(None if blank[i, j] else 0)
        else:
            row.append(int(raw[i, j]))
    ws.append(row)
ss = wb.create_sheet("Samples")
ss.append(["Sample_ID", "Condition"])
perm2 = list(range(12))
random.shuffle(perm2)
for j in perm2:
    ss.append([short[j], cond[j]])
ts = wb.create_sheet("Targets")
ts.append(["Protein_ID"])
for ti in targets:
    ts.append([prot[ti]])
wb.save("raw_intensity.xlsx")
