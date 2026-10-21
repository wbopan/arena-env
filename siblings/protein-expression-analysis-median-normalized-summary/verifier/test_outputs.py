import os
import math
import pandas as pd
from openpyxl import load_workbook

IN = "/root/work/raw_intensity.xlsx"
OUT = "/root/work/target_summary.xlsx"
HDR = ["Protein_ID", "N_Control", "N_Treated", "Mean_Control", "Mean_Treated", "Log2FC", "Fold_Change"]


def expected():
    d = pd.read_excel(IN, sheet_name="Data")
    s = pd.read_excel(IN, sheet_name="Samples")
    t = pd.read_excel(IN, sheet_name="Targets")
    cmap = {str(a).strip(): str(b).strip() for a, b in zip(s["Sample_ID"], s["Condition"])}
    cols = [c for c in d.columns if c != "Protein_ID"]
    norm = {}
    for c in cols:
        vals = {}
        for p, v in zip(d["Protein_ID"], d[c]):
            if v is not None and not (isinstance(v, float) and math.isnan(v)) and float(v) > 0:
                vals[str(p)] = math.log2(float(v))
        sv = sorted(vals.values())
        n = len(sv)
        med = sv[n // 2] if n % 2 else (sv[n // 2 - 1] + sv[n // 2]) / 2
        norm[c] = {k: v - med for k, v in vals.items()}
    group = {c: cmap[c.split("_")[-1]] for c in cols}
    out = []
    for p in t["Protein_ID"]:
        p = str(p)
        cv = [norm[c][p] for c in cols if group[c] == "Control" and p in norm[c]]
        tv = [norm[c][p] for c in cols if group[c] == "Treated" and p in norm[c]]
        mc = sum(cv) / len(cv)
        mt = sum(tv) / len(tv)
        out.append([p, len(cv), len(tv), mc, mt, mt - mc, 2 ** (mt - mc)])
    return out


def rows():
    wb = load_workbook(OUT, data_only=True)
    ws = wb["Summary"]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def test_file_and_sheet():
    assert os.path.exists(OUT)
    wb = load_workbook(OUT)
    assert "Summary" in wb.sheetnames


def test_header_and_protein_order():
    r = rows()
    assert [str(x).strip() if x is not None else None for x in r[0][:7]] == HDR
    exp = expected()
    body = [x for x in r[1:] if x and x[0] is not None]
    assert [str(x[0]).strip() for x in body] == [e[0] for e in exp]


def test_counts():
    exp = expected()
    body = [x for x in rows()[1:] if x and x[0] is not None]
    for b, e in zip(body, exp):
        assert isinstance(b[1], (int, float)) and isinstance(b[2], (int, float))
        assert int(b[1]) == e[1] and int(b[2]) == e[2], b[0]


def test_means_and_log2fc():
    exp = expected()
    body = [x for x in rows()[1:] if x and x[0] is not None]
    for b, e in zip(body, exp):
        for k in (3, 4, 5):
            assert isinstance(b[k], (int, float)), (b[0], k)
            assert abs(b[k] - e[k]) < 2e-4, (b[0], HDR[k], b[k], e[k])


def test_fold_change():
    exp = expected()
    body = [x for x in rows()[1:] if x and x[0] is not None]
    for b, e in zip(body, exp):
        assert isinstance(b[6], (int, float))
        assert abs(b[6] - e[6]) < 2e-4 + 1e-4 * abs(e[6]), (b[0], b[6], e[6])
