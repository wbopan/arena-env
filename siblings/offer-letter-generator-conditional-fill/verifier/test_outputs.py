import json
import re
from pathlib import Path

import pytest
from docx import Document

WORK = Path("/root/work")
TEMPLATE = WORK / "offer_letter_template.docx"
DATA = WORK / "employee_data.json"
OUT = WORK / "offer_letter_filled.docx"
PH = re.compile(r"\{\{([A-Z_]+)\}\}")
MARK = re.compile(r"^\{\{(END_)?IF_([A-Z_]+)\}\}$")


def table_texts(doc):
    out = []

    def walk(t):
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    out.append(p.text)
                for nt in cell.tables:
                    walk(nt)

    for t in doc.tables:
        walk(t)
    return out


def hf_texts(doc):
    out = []
    for s in doc.sections:
        out += [p.text for p in s.header.paragraphs]
        out += [p.text for p in s.footer.paragraphs]
    return out


def all_text(doc):
    return "\n".join([p.text for p in doc.paragraphs] + table_texts(doc) + hf_texts(doc))


def sub(text, fields):
    return PH.sub(lambda m: str(fields[m.group(1)]), text).strip()


def body_blocks(tmpl):
    cur = None
    res = []
    for p in tmpl.paragraphs:
        m = MARK.match(p.text.strip())
        if m:
            cur = None if m.group(1) else m.group(2)
            continue
        res.append((cur, p.text))
    return res


@pytest.fixture(scope="module")
def data():
    return json.loads(DATA.read_text())


@pytest.fixture(scope="module")
def tmpl():
    return Document(str(TEMPLATE))


@pytest.fixture(scope="module")
def out():
    assert OUT.exists(), f"missing {OUT}"
    return Document(str(OUT))


def test_output_is_valid_docx(out):
    assert len(out.paragraphs) > 0


def test_no_leftover_braces(out):
    txt = all_text(out)
    assert "{{" not in txt and "}}" not in txt, re.findall(r"\{\{[^}]*\}\}", txt)


def test_body_paragraphs(out, tmpl, data):
    body = {p.text.strip() for p in out.paragraphs}
    for cond, text in body_blocks(tmpl):
        if cond is None and PH.search(text):
            exp = sub(text, data["fields"])
            assert exp in body, f"expected body paragraph: {exp!r}"


def test_tables_including_nested(out, tmpl, data):
    got = {t.strip() for t in table_texts(out)}
    for text in table_texts(tmpl):
        if PH.search(text):
            exp = sub(text, data["fields"])
            assert exp in got, f"expected table text: {exp!r}"


def test_header_footer(out, tmpl, data):
    got = {t.strip() for t in hf_texts(out)}
    for text in hf_texts(tmpl):
        if PH.search(text):
            exp = sub(text, data["fields"])
            assert exp in got, f"expected header/footer text: {exp!r}"


def test_conditional_blocks(out, tmpl, data):
    body = {p.text.strip() for p in out.paragraphs}
    txt = all_text(out)
    for cond, text in body_blocks(tmpl):
        if cond is None:
            continue
        exp = sub(text, data["fields"])
        if data["conditions"][cond] == "Yes":
            assert exp in body, f"kept block content missing: {exp!r}"
        else:
            assert exp not in txt, f"excluded block content present: {exp!r}"
            lit = max(re.split(r"\{\{[A-Z_]+\}\}", text), key=len).strip()
            if len(lit) >= 10:
                assert lit not in txt, f"excluded text present: {lit!r}"
