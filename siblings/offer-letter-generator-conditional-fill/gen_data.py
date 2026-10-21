import json
import random
from datetime import date, timedelta

from docx import Document

random.seed(731)

firsts = ["Priya", "Marcus", "Elena", "Tomasz", "Aisha", "Kenji", "Sofia", "Daniel"]
lasts = ["Okafor", "Lindqvist", "Haddad", "Moreau", "Castillo", "Nakamura", "Brennan"]
staff = ["Grace Whitfield", "Omar Benali", "Hannah Iverson", "Luis Ortega", "Mei Tanaka"]
mgr, hr = random.sample(staff, 2)
base = date(2024, 3, 1) + timedelta(days=random.randint(0, 200))
fields = {
    "COMPANY_NAME": random.choice(["Northwind Analytics", "Bluepeak Systems", "Cedar Ridge Labs"]),
    "CANDIDATE_FULL_NAME": random.choice(firsts) + " " + random.choice(lasts),
    "POSITION": random.choice(["Senior Data Engineer", "Product Manager", "Financial Analyst", "UX Researcher"]),
    "DEPARTMENT": random.choice(["Platform", "Finance", "Growth", "Research"]),
    "DATE": base.strftime("%B %d, %Y"),
    "RESPONSE_DEADLINE": (base + timedelta(days=7)).strftime("%B %d, %Y"),
    "START_DATE": (base + timedelta(days=30)).strftime("%B %d, %Y"),
    "MANAGER_NAME": mgr,
    "HR_NAME": hr,
    "BASE_SALARY": "${:,}".format(random.randrange(85000, 190000, 500)),
    "PTO_DAYS": str(random.randint(15, 30)),
    "RELOCATION_AMOUNT": "${:,}".format(random.randrange(5000, 20000, 250)),
    "SIGNING_BONUS": "${:,}".format(random.randrange(3000, 25000, 500)),
    "CITY": random.choice(["Austin", "Denver", "Seattle", "Raleigh", "Boston"]),
    "OFFER_ID": "OFR-" + str(random.randint(10000, 99999)),
}
reloc = random.random() < 0.5
conditions = {"RELOCATION": "Yes" if reloc else "No", "SIGNING_BONUS": "No" if reloc else "Yes"}
with open("employee_data.json", "w") as f:
    json.dump({"fields": fields, "conditions": conditions}, f, indent=2)


def runs(p, parts):
    for t in parts:
        p.add_run(t)
    return p


doc = Document()
sec = doc.sections[0]
runs(sec.header.paragraphs[0], ["Confidential - ", "{{COMPANY_", "NAME}}", " offer for {{CANDIDATE_", "FULL_NAME}}"])
runs(sec.footer.paragraphs[0], ["Ref: {{OFFER_", "ID}} | Contact: {{HR_", "NAME}}"])
body = [
    ["{{DA", "TE}}"],
    ["Dear {{CANDIDATE_", "FULL_NAME}},"],
    ["We are pleased to offer you the position of ", "{{POSI", "TION}}", " in the ", "{{DEPART", "MENT}}", " department, starting ", "{{START_", "DATE}}", "."],
]
for parts in body:
    runs(doc.add_paragraph(), parts)
t = doc.add_table(rows=2, cols=2)
t.cell(0, 0).text = "Compensation summary"
runs(t.cell(0, 1).paragraphs[0], ["Offer ", "{{OFFER_", "ID}}"])
runs(t.cell(1, 0).paragraphs[0], ["Reports to: {{MANAGER_", "NAME}}"])
nt = t.cell(1, 1).add_table(rows=2, cols=2)
nt.cell(0, 0).text = "Base salary"
runs(nt.cell(0, 1).paragraphs[0], ["{{BASE_", "SALARY}}", " per year"])
nt.cell(1, 0).text = "Paid time off"
runs(nt.cell(1, 1).paragraphs[0], ["{{PTO_", "DAYS}} days"])
rest = [
    ["{{IF_", "RELOCATION}}"],
    ["Relocation assistance: we will provide ", "{{RELOCATION_", "AMOUNT}}", " to support your move to ", "{{CI", "TY}}", "."],
    ["Temporary corporate housing is included for 30 days."],
    ["{{END_IF_", "RELOCATION}}"],
    ["{{IF_SIGNING", "_BONUS}}"],
    ["Signing bonus: you will receive a one-time payment of ", "{{SIGNING_", "BONUS}}", " with your first paycheck."],
    ["{{END_IF_", "SIGNING_BONUS}}"],
    ["Please respond by ", "{{RESPONSE_", "DEADLINE}}", "."],
    ["Sincerely, ", "{{HR_", "NAME}}"],
]
for parts in rest:
    runs(doc.add_paragraph(), parts)
doc.save("offer_letter_template.docx")
